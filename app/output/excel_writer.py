"""Excel 导出：原始数据 / 情感编码明细 / 统计汇总 / 关键词效果 / 实际查询串。"""

from __future__ import annotations

from io import BytesIO
import json
import re

import pandas as pd

from app.core.models import ReportBundle
from app.core.names import ATTRIBUTION_CN, NARRATIVE_CN, dimension_cn
from app.coding.coder import display_confidence, display_confidence_tier
from app.core.names import register_custom_dim_names
from app.domains.loader import task_schema
from app.output.html_report import keyword_rows, query_rows

SENTIMENT_CN = {
    "positive": "正面",
    "negative": "负面",
    "neutral": "中性",
}

PLATFORM_CN = {
    "websearch": "全网搜索",
    "websearch_zhihu": "知乎",
    "websearch_tieba": "贴吧",
    "websearch_taptap": "TapTap",
    "bilibili": "B站",
    "weibo": "微博",
    "xiaohongshu": "小红书",
}

INTENSITY_CN = {
    1: "很弱",
    2: "较弱",
    3: "一般",
    4: "较强",
    5: "很强",
}


def _posts_rows(bundle: ReportBundle) -> list[dict]:
    rows = []
    for result in bundle.channel_results:
        for post in result.posts:
            spec = post.platform_specific or {}
            rows.append(
                {
                    "平台": post.platform,
                    "关键词": post.keyword,
                    "作者": post.author,
                    "标题": post.title,
                    "正文": post.content,
                    "链接": post.url,
                    "发布时间": post.timestamp,
                    "点赞数": post.likes,
                    "转发数": post.reposts,
                    "评论数": post.comments_count,
                    "已采集评论数": len(post.comments),
                    "评论内容": " || ".join(c.text for c in post.comments),
                    "平台特有字段": "；".join(f"{k}={v}" for k, v in spec.items()),
                }
            )
    return rows


def _comments_rows(bundle: ReportBundle) -> list[dict]:
    """评论明细：每条评论一行，避免在单元格里被截断。"""
    rows = []
    for result in bundle.channel_results:
        for post in result.posts:
            for comment in post.comments:
                rows.append(
                    {
                        "帖子链接": post.url,
                        "平台": post.platform,
                        "关键词": post.keyword,
                        "帖子作者": post.author,
                        "评论作者": comment.author,
                        "评论内容": comment.text,
                        "点赞数": comment.likes,
                        "评论时间": comment.time,
                        "是否回复": "是（楼中楼）" if comment.is_reply else "否",
                    }
                )
    return rows


def _simple_rows(bundle: ReportBundle) -> list[dict]:
    """小白视图：字段全部用中文人话，去掉内部加工痕迹，保留可溯源内容。"""
    dim_names: dict[str, str] = {}
    register_custom_dim_names(bundle.plan)
    schema = task_schema(bundle.plan)
    if schema:
        dim_names = {d.id: d.name for d in schema.dimensions}

    def _norm(t: str) -> str:
        # 编码管道会把换行/空格规整为逗号，映射时按同样规则归一
        t = re.sub(r"\s+", ",", (t or "").strip())
        return t.strip("，, ")

    # 原文链接：评论按内容（原始/规整）匹配，帖子按 text_id 里的 URL 片段匹配
    post_urls: list[str] = []
    comment_urls: dict[str, str] = {}
    for ch in bundle.channel_results:
        for post in ch.posts:
            post_urls.append(post.url)
            for c in post.comments:
                ct = (c.text or "").strip()
                comment_urls[ct] = post.url
                comment_urls[_norm(ct)] = post.url
    rows = []
    for it in bundle.coded_items:
        sentiment = SENTIMENT_CN.get(it.sentiment.value, it.sentiment.value)
        reason = ""
        text = (it.text or "").strip()
        if sentiment == "负面":
            if any(k in text for k in ("追责", "恢复", "下架", "要求", "投诉", "抵制", "退钱")):
                reason = "投诉/诉求"
            elif any(k in text for k in ("像", "越来越", "呵呵", "呵呵哒", "真好", "果然", "太棒了")):
                reason = "反讽/吐槽"
            else:
                reason = "负面评价"
        elif sentiment == "正面":
            reason = "正面评价"
        else:
            if any(k in text for k in ("?", "？", "怎么", "如何", "选", "推荐", "对比")):
                reason = "咨询/提问"
            elif not text:
                reason = "正文为空"
            else:
                reason = "无明显倾向"
        # 判断来源：评论优先按内容匹配，帖子按 text_id 里的 URL 片段匹配
        link = comment_urls.get(text) or comment_urls.get(_norm(text)) or ""
        if not link:
            tid_key = it.text_id.rsplit(":", 1)[0]  # 去掉 :post/:comment 后缀
            for u in post_urls:
                if tid_key and (tid_key in u or u.endswith(tid_key)):
                    link = u
                    break
        rows.append(
            {
                "文本内容": text,
                "平台": PLATFORM_CN.get(it.platform, it.platform),
                "日期": it.pub_date,
                "大家怎么说": sentiment,
                "判定依据": reason,
                "可信度": display_confidence_tier(it.confidence, it.method),
                "提到什么": "、".join(it.keywords),
                "维度": "、".join(dim_names.get(d, d) for d in it.dimensions),
                "维度情感": "、".join(
                    f"{dim_names.get(d, d)}：{SENTIMENT_CN.get(v, v)}"
                    for d, v in it.dimension_sentiments.items()
                ) or "无",
                "来源渠道": PLATFORM_CN.get(it.platform, it.platform),
                "原文链接": link,
            }
        )
    return rows


def _coded_rows(bundle: ReportBundle) -> list[dict]:
    dim_names: dict[str, str] = {}
    register_custom_dim_names(bundle.plan)
    schema = task_schema(bundle.plan)
    if schema:
        dim_names = {d.id: d.name for d in schema.dimensions}

    def _note(it) -> str:
        notes = []
        if it.need_review:
            notes.append(f"需复核：{it.need_review_reason}")
            if not it.dimensions:
                notes.append("维度：未命中领域词表")
            if not bundle.plan.narrative_enabled:
                notes.append("叙事/归因：未启用")
            elif it.method != "llm":
                notes.append("叙事/归因：仅对 LLM 精分析文本执行")
            elif not it.narrative and not it.attribution:
                notes.append("叙事/归因：模型未给出")
            return "；".join(notes)
        return ""

    return [
        {
            "文本ID": it.text_id,
            "文本内容": it.text,
            "平台": it.platform,
            "关键词": it.keyword,
            "发布日期": it.pub_date,
            "情感": it.sentiment.value,
            "情感评分": it.sentiment_score,
            "置信度": display_confidence(it.confidence, it.method),
            "强度(1-5)": it.intensity,
            "分析方法": "词典直判" if it.method == "lexicon" else "LLM 精分析",
            "维度": "、".join(dim_names.get(d, d) for d in it.dimensions),
            "维度情感": json.dumps(
                {dim_names.get(d, d): v for d, v in it.dimension_sentiments.items()},
                ensure_ascii=False,
            ) or "{}",
            "情感关键词": "、".join(it.keywords),
            "叙事框架": NARRATIVE_CN.get(it.narrative.value, it.narrative.value)
            if it.narrative
            else "",
            "归因主体": ATTRIBUTION_CN.get(it.attribution, it.attribution) if it.attribution else "",
            "需复核": "是" if it.need_review else "",
            "说明": _note(it),
        }
        for it in bundle.coded_items
    ]


def _dropped_rows(bundle: ReportBundle) -> list[dict]:
    rows = []
    for result in bundle.channel_results:
        for d in result.dropped:
            rows.append(
                {
                    "平台": d.get("platform", ""),
                    "链接": d.get("url", ""),
                    "标题": d.get("title", ""),
                    "关键词": d.get("keyword", ""),
                    "查询串": d.get("query", ""),
                    "丢弃原因": d.get("reason", ""),
                    "正文摘要": d.get("content", ""),
                    "判定依据": d.get("match", ""),
                }
            )
    return rows


def _summary_frames(bundle: ReportBundle) -> list[tuple[str, pd.DataFrame]]:
    s = bundle.summary
    frames: list[tuple[str, pd.DataFrame]] = []
    collected = sum(
        len(ch.posts) + len(ch.dropped)
        for ch in bundle.channel_results
        if ch.ok
    )
    kept = sum(len(ch.posts) for ch in bundle.channel_results)
    dropped_n = sum(len(ch.dropped) for ch in bundle.channel_results)
    comment_n = sum(
        1 for ch in bundle.channel_results for p in ch.posts for _ in p.comments
    )
    frames.append(
        (
            "采集漏斗",
            pd.DataFrame(
                [
                    {"阶段": "渠道采集帖子", "数量": collected},
                    {"阶段": "清洗保留（丢弃 %d）" % dropped_n, "数量": kept},
                    {"阶段": "评论（热门+楼中楼）", "数量": comment_n},
                    {"阶段": "情感编码文本（正文+评论）", "数量": s["total_items"]},
                ]
            ),
        )
    )
    frames.append(
        (
            "字段说明",
            pd.DataFrame(
                [
                    {"字段": "维度", "说明": "命中领域 schema 关键词的维度提及；未命中留空并在明细说明列注明"},
                    {"字段": "维度情感", "说明": "2.4 维度级情感：仅对明确带褒贬的维度标注正面/负面（转折句逐维拆解）；无明确褒贬留空；词典模式为轻量兜底"},
                    {"字段": "叙事框架/归因主体", "说明": "仅对 LLM 精分析过的文本执行（成本控制设计）"},
                    {"字段": "已采集评论数", "说明": "原始数据中的平台评论总数来自平台字段；实际采集按每帖上限抓取热门评论"},
                    {"字段": "丢弃明细", "说明": "清洗阶段被排除的帖子及原因（官方页面/样板文本/重复/不相关等）"},
                    {"字段": "关键词效果", "说明": "按确认关键词统计采集/保留/丢弃/有效供给率（保留÷采集），与 HTML/Word 报告口径一致"},
                    {"字段": "实际查询串（按渠道）", "说明": "系统实际发给各渠道的查询词（WebSearch：确认词+子渠道提示+自动后缀；B站/微博/小红书：策略展开）；可核对'实际搜了什么'"},
                ]
            ),
        )
    )
    kw_rows, kw_unattr = keyword_rows(bundle)
    if kw_rows:
        frames.append(
            (
                "关键词效果",
                pd.DataFrame(
                    [
                        {
                            "关键词": r["keyword"],
                            "采集": r["collected"],
                            "保留": r["kept"],
                            "丢弃": r["dropped"],
                            "有效供给率": r["effective_rate"],
                            "编码文本": r["coded"],
                            "正面": r["positive"],
                            "负面": r["negative"],
                            "中性": r["neutral"],
                            "负面率": r["negative_rate"],
                        }
                        for r in kw_rows
                    ]
                ),
            )
        )
    q_rows, q_unattr = query_rows(bundle)
    if q_rows:
        frames.append(
            (
                "实际查询串（按渠道）",
                pd.DataFrame(
                    [
                        {
                            "实际查询串": r["query"],
                            "渠道": r["channel"],
                            "采集": r["collected"],
                            "保留": r["kept"],
                            "丢弃": r["dropped"],
                            "有效供给率": r["effective_rate"],
                            "编码文本": r["coded"],
                            "负面": r["negative"],
                            "负面率": r["negative_rate"],
                        }
                        for r in q_rows
                    ]
                ),
            )
        )
    if kw_unattr or q_unattr:
        for idx, (name, df) in enumerate(frames):
            if name == "字段说明":
                frames[idx][1].loc[len(df)] = {
                    "字段": "未归属丢弃",
                    "说明": f"旧版数据未记录关键词/查询串的丢弃 {kw_unattr + q_unattr} 条，不计入上表",
                }
                break
    usage = bundle.llm_usage or {}
    if usage.get("prompt_tokens"):
        frames.append(
            (
                "LLM 用量与费用",
                pd.DataFrame(
                    [
                        {"指标": "输入 token", "数值": usage.get("prompt_tokens")},
                        {"指标": "输出 token", "数值": usage.get("completion_tokens")},
                        {
                            "指标": "预估费用（元）",
                            "数值": usage.get("estimated_cost"),
                        },
                        {
                            "指标": "说明",
                            "数值": "以 DeepSeek 官方计费为准，单价可能调整",
                        },
                    ]
                ),
            )
        )
    frames.append(
        (
            "整体概览",
            pd.DataFrame(
                [
                    {"指标": "分析对象", "数值": bundle.plan.subject},
                    {"指标": "渠道数", "数值": len(bundle.channel_results)},
                    {"指标": "帖子数", "数值": s["total_posts"]},
                    {"指标": "编码文本数", "数值": s["total_items"]},
                    {"指标": "广告/官方内容",
                     "数值": (f"{s.get('ads', {}).get('count', 0)} 条"
                              f"（占 {s.get('ads', {}).get('ratio_of_total', 0):.1%}，"
                              f"{s.get('ads', {}).get('mode', '计入')}）")},
                    {"指标": "整体倾向", "数值": s["overall_sentiment"]},
                    {"指标": "平均情感分", "数值": s["avg_score"]},
                ]
            ),
        )
    )
    dist = pd.DataFrame(
        [
            {
                "情感": name,
                "数量": s["sentiment_distribution"][key]["count"],
                "占比": s["sentiment_distribution"][key]["ratio"],
            }
            for key, name in [("positive", "正面"), ("negative", "负面"), ("neutral", "中性")]
        ]
    )
    frames.append(("情感分布", dist))
    frames.append(
        (
            "平台统计",
            pd.DataFrame(
                [
                    {
                        "平台": pid,
                        "内容数": v["posts"],
                        "平均评分": v["avg_score"],
                        "正面": v["positive"],
                        "负面": v["negative"],
                        "中性": v["neutral"],
                    }
                    for pid, v in s["platforms"].items()
                ]
            ),
        )
    )
    frames.append(
        (
            "维度统计",
            pd.DataFrame(
                [
                    {
                        "维度": dimension_cn(did),
                        "评价量": v["count"],
                        "负面数": v["negative"],
                        "负面率": v["negative_rate"],
                    }
                    for did, v in s["dimensions"].items()
                ]
            ),
        )
    )
    frames.append(
        (
            "时间趋势",
            pd.DataFrame(
                [
                    {"日期": d, "内容量": v["count"], "平均评分": v["avg_score"], "负面数": v["negative"]}
                    for d, v in s["trend"].items()
                ]
            ),
        )
    )
    return frames


def _apply_sentiment_colors(ws) -> None:
    """给"大家怎么说"列加色块：正面绿、负面红、中性灰。"""
    from openpyxl.styles import PatternFill

    fill_map = {
        "正面": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        "负面": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        "中性": PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid"),
    }
    header_row = [c.value for c in ws[1]]
    if "大家怎么说" not in header_row:
        return
    col_idx = header_row.index("大家怎么说") + 1
    for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
        cell = row[0]
        fill = fill_map.get(str(cell.value or "").strip())
        if fill:
            cell.fill = fill


def build_excel(bundle: ReportBundle) -> BytesIO:
    """生成 Excel 文件字节流（小白视图 + 原始数据 + 评论明细 + 丢弃明细 + 情感编码明细 + 统计汇总）。"""
    out = BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        simple = pd.DataFrame(_simple_rows(bundle))
        simple.to_excel(writer, sheet_name="小白视图", index=False)
        _apply_sentiment_colors(writer.sheets["小白视图"])

        raw = pd.DataFrame(_posts_rows(bundle))
        raw.to_excel(writer, sheet_name="原始数据", index=False)

        comments = pd.DataFrame(_comments_rows(bundle))
        comments.to_excel(writer, sheet_name="评论明细", index=False)

        dropped = pd.DataFrame(_dropped_rows(bundle))
        dropped.to_excel(writer, sheet_name="丢弃明细", index=False)

        coded = pd.DataFrame(_coded_rows(bundle))
        coded.to_excel(writer, sheet_name="情感编码明细", index=False)

        summary_rows = []
        for name, df in _summary_frames(bundle):
            summary_rows.append(f"【{name}】")
            summary_rows.extend(df.to_string(index=False).splitlines())
        pd.DataFrame({"统计汇总": summary_rows}).to_excel(
            writer, sheet_name="统计汇总", index=False
        )

        # F-010（2026-08-26）：解读与建议 sheet（与 HTML/Word 同源：structured_summary）
        if bundle.structured_summary:
            _ss = bundle.structured_summary
            _ss_source = getattr(bundle, "structured_summary_source", "rule")
            _ss_rows: list[dict[str, str]] = [
                {
                    "项目": "来源",
                    "内容": "AI 归因（LLM 解读）" if _ss_source == "llm" else "规则推测（非 AI 归因）",
                },
                {"项目": "一句话结论", "内容": _ss.get("overall", "")},
            ]
            _pos = _ss.get("positive") or {}
            _neg = _ss.get("negative") or {}
            if _pos:
                _ss_rows.append(
                    {"项目": "正面反馈占比", "内容": f"{_pos.get('ratio', 0) * 100:.1f}%"}
                )
                for _p in _pos.get("phrases") or []:
                    _ss_rows.append(
                        {
                            "项目": "正面代表短语",
                            "内容": (
                                f"{_p.get('phrase', '')}"
                                f"（{_p.get('count', 0)} 条/{_p.get('ratio', 0) * 100:.0f}%）"
                            ),
                        }
                    )
            if _neg:
                _ss_rows.append(
                    {"项目": "负面反馈占比", "内容": f"{_neg.get('ratio', 0) * 100:.1f}%"}
                )
                for _p in _neg.get("phrases") or []:
                    _ss_rows.append(
                        {
                            "项目": "负面代表短语",
                            "内容": (
                                f"{_p.get('phrase', '')}"
                                f"（{_p.get('count', 0)} 条/{_p.get('ratio', 0) * 100:.0f}%）"
                            ),
                        }
                    )
            for _issue in _ss.get("top_issues") or []:
                _ss_rows.append(
                    {
                        "项目": f"重点问题：{_issue.get('name', '')}",
                        "内容": (
                            f"负面率 {_issue.get('rate', 0) * 100:.0f}%"
                            f"（n={_issue.get('count', 0)}）"
                        ),
                    }
                )
                if _issue.get("cause"):
                    _ss_rows.append({"项目": "可能原因（规则推测）", "内容": _issue["cause"]})
                if _issue.get("direction"):
                    _ss_rows.append({"项目": "建议", "内容": _issue["direction"]})
            for _im in _ss.get("improvements") or []:
                _ss_rows.append({"项目": "改进建议", "内容": _im})
            pd.DataFrame(_ss_rows).to_excel(
                writer, sheet_name="解读与建议", index=False
            )

        # F-015 P2（2026-08-26）：主题洞察 sheet（短语归并卡）
        _topics = (bundle.summary or {}).get("topics") or []
        if _topics:
            _topic_rows: list[dict[str, str]] = []
            for _t in _topics[:8]:
                _tpol = {"positive": "正面", "negative": "负面", "neutral": "中性"}.get(
                    _t.get("polarity", ""), "中性"
                )
                _topic_rows.append({
                    "主题": _t.get("name", ""),
                    "维度": dimension_cn(_t.get("dimension", "")),
                    "提及量": str(_t.get("count", 0)),
                    "主导情感": _tpol,
                    "代表短语": "、".join(_t.get("phrases") or []),
                })
            pd.DataFrame(_topic_rows).to_excel(
                writer, sheet_name="主题洞察", index=False
            )
    out.seek(0)
    return out
