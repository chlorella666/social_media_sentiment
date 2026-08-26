"""Word 静态报告（python-docx）。

MVP 输出文本与统计表格；V1 将图表渲染为 PNG 后嵌入。
"""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt

from app.core.models import ReportBundle
from app.output.html_report import (
    cooccurrence_fig,
    cooccurrence_plan,
    date_dim_heatmap_fig,
    dimensions_fig,
    heatmap_fig,
    intensity_fig,
    keyword_rows,
    narrative_actor_fig,
    narrative_frame_actor_heatmap,
    narrative_insight_text,
    overall_fig,
    platform_dim_fig,
    platform_fig,
    query_rows,
    radar_fig,
    sentiment_sources_fig,
    topic_cluster_rows,
    topic_pairs,
    trend_fig,
    wordcloud_png_bytes,
    words_fig,
)
from app.core.names import dimension_cn
from app.core.evidence import (
    dimension_evidence_label,
    display_action,
    display_finding_id,
    findings_section_title,
)

SENTIMENT_NAMES = {"positive": "正面", "negative": "负面", "neutral": "中性"}

CHART_BUILDERS = [
    ("overall", "整体情感占比", overall_fig),
    ("platform", "各平台情感分布", platform_fig),
    ("trend", "情感趋势（内容量 + 平均评分 + 负面率）", trend_fig),
    ("dimensions", "维度讨论量与负面率", dimensions_fig),
    ("heatmap", "维度负面率热力图", heatmap_fig),
    ("intensity", "情绪强度直方图", intensity_fig),
    ("radar", "各维度负面率/平均分雷达图", radar_fig),
    ("platform_dim", "平台 × 维度负面率", platform_dim_fig),
    ("date_dim", "日期 × 维度负面率热力图", date_dim_heatmap_fig),
    ("words", "代表观点（短语）Top20", words_fig),
    ("sources", "负面情绪来源话题榜", sentiment_sources_fig),
    ("cooccurrence", "关键词共现网络图", cooccurrence_fig),
    ("narrative_actor", "归因主体分布（用户主要把问题归给谁）", narrative_actor_fig),
    ("narrative_heatmap", "叙事框架 × 归因主体负面率", narrative_frame_actor_heatmap),
]

_calc_fig = None


def _get_calc_fig():
    """kaleido v1：启动一次渲染服务，之后所有图表共用（大幅提速）。"""
    global _calc_fig
    if _calc_fig is None:
        try:
            from kaleido import calc_fig_sync, start_sync_server

            start_sync_server(silence_warnings=True)
            _calc_fig = calc_fig_sync
        except Exception:
            _calc_fig = False
    return _calc_fig or None


def _add_chart_image(doc: Document, fig, title: str) -> None:
    """用 kaleido 把 plotly 图渲染成 PNG 嵌入 Word。"""
    calc = _get_calc_fig()
    if calc is None:
        return
    try:
        data = calc(
            fig,
            opts=dict(format="png", width=850, height=480, scale=1.1),
        )
    except Exception:
        return
    doc.add_picture(BytesIO(data), width=Inches(6.2))
    run = doc.add_paragraph().add_run(title)
    run.bold = True


def _append_keyword_appendix(doc: Document, bundle: ReportBundle) -> None:
    """附录：关键词效果与采集明细（2026-08-19：从正文前段移入文末，
    与 HTML「方法与数据说明 → 附录」结构对齐）。"""
    doc.add_heading("十、附录：关键词效果与采集明细", level=1)
    doc.add_paragraph(
        "以下为方法与数据说明：记录系统实际搜了什么、每个词/查询串采了多少、"
        "留了多少、丢了多少。"
    )
    kw_rows, unattr = keyword_rows(bundle)
    if kw_rows:
        ktable = doc.add_table(rows=1, cols=10)
        ktable.style = "Table Grid"
        hdr = ktable.rows[0].cells
        for i, name in enumerate(
            ["关键词", "采集", "保留", "丢弃", "有效供给率",
             "编码文本", "正面", "负面", "中性", "负面率"]
        ):
            hdr[i].text = name
        for r in kw_rows:
            cells = ktable.add_row().cells
            for i, key in enumerate(
                ["keyword", "collected", "kept", "dropped", "effective_rate",
                 "coded", "positive", "negative", "neutral", "negative_rate"]
            ):
                cells[i].text = str(r[key])
        if unattr:
            doc.add_paragraph(
                f"另有 {unattr} 条丢弃记录未归属到关键词（旧版数据），不计入上表。"
            )
        doc.add_paragraph(
            "说明：采集=该关键词采回的帖子数（含清洗丢弃）；"
            "有效供给率=保留/采集；样本少的行仅供参考。"
        )
    else:
        doc.add_paragraph("（本次任务无关键词级统计）")

    q_rows, q_unattr = query_rows(bundle)
    if q_rows:
        p = doc.add_paragraph()
        run = p.add_run("实际查询串（按渠道）")
        run.bold = True
        doc.add_paragraph(
            "说明：实际查询串是系统实际发给各渠道的查询词（WebSearch 可能在确认"
            "关键词基础上追加后缀/子渠道提示/策略词；B站/微博/小红书启用渠道关键词"
            "优化时按策略展开）；同一确认关键词可能对应多个实际查询串。"
        )
        qtable = doc.add_table(rows=1, cols=9)
        qtable.style = "Table Grid"
        hdr = qtable.rows[0].cells
        for i, name in enumerate(
            ["实际查询串", "渠道", "采集", "保留", "丢弃",
             "有效供给率", "编码文本", "负面", "负面率"]
        ):
            hdr[i].text = name
        for r in q_rows:
            cells = qtable.add_row().cells
            for i, key in enumerate(
                ["query", "channel", "collected", "kept", "dropped",
                 "effective_rate", "coded", "negative", "negative_rate"]
            ):
                cells[i].text = str(r[key])
        if q_unattr:
            doc.add_paragraph(
                f"另有 {q_unattr} 条 WebSearch 丢弃记录未归属到查询串（旧版数据），不计入上表。"
            )


def build_word(bundle: ReportBundle) -> BytesIO:
    doc = Document()
    doc.add_heading(f"社交媒体情感分析报告 — {bundle.plan.subject}", level=0)
    doc.add_paragraph(f"生成时间：{bundle.created_at.strftime('%Y-%m-%d %H:%M')}")
    if bundle.insight_mode == "lexicon":
        doc.add_paragraph(
            "⚠ 本次为词典模式，维度情感与结论仅供参考；"
            "开启 LLM 精分析可获得可归因的结论与行动建议。"
        )
    elif bundle.insight_mode == "template_fallback":
        doc.add_paragraph(
            "⚠ 本次深度结论由规则模板生成（LLM 兜底），仅供参考；"
            "建议检查 LLM 配置后重新生成。"
        )
    elif bundle.insight_mode == "no_data":
        doc.add_paragraph("⚠ 本次未采集到有效文本，未生成情感结论。")
    elif bundle.insight_mode == "review_refresh":
        doc.add_paragraph(
            "⚠ 报告已按人工复核结果用规则重新生成，结论仅供参考；"
            "如需 LLM 深度结论，请重新生成。"
        )

    s = bundle.summary
    doc.add_heading("一、分析概览", level=1)
    overview = doc.add_table(rows=0, cols=2)
    overview.style = "Table Grid"
    for k, v in [
        ("分析对象", bundle.plan.subject),
        ("关键词数", str(len(bundle.plan.keywords))),
        ("帖子数", str(s["total_posts"])),
        ("编码文本数", str(s["total_items"])),
        ("整体倾向", s["overall_sentiment"]),
        ("平均情感分", str(s["avg_score"])),
    ]:
        cells = overview.add_row().cells
        cells[0].text, cells[1].text = k, v

    doc.add_heading("二、情感分布", level=1)
    dist = s["sentiment_distribution"]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "情感", "数量", "占比"
    for key, name in [("positive", "正面"), ("negative", "负面"), ("neutral", "中性")]:
        row = table.add_row().cells
        row[0].text = name
        row[1].text = str(dist[key]["count"])
        row[2].text = f"{dist[key]['ratio'] * 100:.1f}%"

    doc.add_heading("三、平台统计", level=1)
    ptable = doc.add_table(rows=1, cols=6)
    ptable.style = "Table Grid"
    hdr = ptable.rows[0].cells
    for i, name in enumerate(["平台", "内容数", "平均评分", "正面", "负面", "中性"]):
        hdr[i].text = name
    for pid, v in s["platforms"].items():
        row = ptable.add_row().cells
        row[0].text = pid
        row[1].text = str(v["posts"])
        row[2].text = str(v["avg_score"])
        row[3].text = str(v["positive"])
        row[4].text = str(v["negative"])
        row[5].text = str(v["neutral"])

    if s["dimensions"]:
        doc.add_heading("四、维度分析", level=1)
        dtable = doc.add_table(rows=1, cols=4)
        dtable.style = "Table Grid"
        hdr = dtable.rows[0].cells
        for i, name in enumerate(["维度", "评价量", "负面数", "负面率"]):
            hdr[i].text = name
        for did, v in s["dimensions"].items():
            row = dtable.add_row().cells
            row[0].text = dimension_cn(did)
            row[1].text = str(v["count"])
            row[2].text = str(v["negative"])
            row[3].text = f"{v['negative_rate'] * 100:.1f}%"
        if sum(1 for it in bundle.coded_items if it.method == "llm") == 0:
            doc.add_paragraph(
                "注：本次为词典模式，维度情感为词典轻量估算，仅供参考；"
                "开启 LLM 精分析后维度情感更准确。"
            )
        neg_top = [c for c in bundle.evidence if c.get("dimension") and c["sentiment"] == "negative"]
        if neg_top:
            p = doc.add_paragraph()
            run = p.add_run("各维度负面原文 Top 3（规则抽取）")
            run.bold = True
            by_dim: dict[str, list[dict]] = {}
            for c in neg_top:
                by_dim.setdefault(c["dimension"], []).append(c)
            for dim, cards in by_dim.items():
                p = doc.add_paragraph()
                run = p.add_run(
                    f"{dimension_cn(dim)}（{dimension_evidence_label(cards)}）"
                )
                run.bold = True
                for c in cards:
                    label = " · 待复核" if c.get("need_review") else ""
                    doc.add_paragraph(
                        f"「{c['text']}」\n"
                        f"来源：{c['platform']} · {c.get('date') or '日期未知'}{label}"
                    )

    doc.add_heading("五、图表与解析", level=1)
    s = bundle.summary
    for cid, title, builder in CHART_BUILDERS:
        fig = builder(s)
        if fig is None:
            continue
        _add_chart_image(doc, fig, title)
        insight = bundle.chart_insights.get(cid, "")
        if insight:
            p = doc.add_paragraph(insight)
            p.paragraph_format.first_line_indent = Pt(24)
    narr_lines = narrative_insight_text(s)
    if narr_lines:
        p = doc.add_paragraph()
        run = p.add_run("叙事/归因解读")
        run.bold = True
        for line in narr_lines:
            doc.add_paragraph(line)
    cluster_rows = topic_cluster_rows(s)
    if cluster_rows:
        p = doc.add_paragraph()
        run = p.add_run("话题簇榜单")
        run.bold = True
        ctable = doc.add_table(rows=1, cols=4)
        ctable.style = "Table Grid"
        hdr = ctable.rows[0].cells
        for i, name in enumerate(["簇名", "代表短语", "涉及文本数", "负面率"]):
            hdr[i].text = name
        for r in cluster_rows:
            cells = ctable.add_row().cells
            cells[0].text = r["name"]
            cells[1].text = r["words"]
            cells[2].text = str(r["doc_count"])
            cells[3].text = r["negative_rate"]
    co_plan = cooccurrence_plan(s)
    if co_plan["kind"] == "pairs":
        doc.add_paragraph(
            f"注：讨论结构样本不足，已显示话题词对榜（{co_plan['reason']}）。"
        )
        pair_rows = topic_pairs(s)
        if pair_rows:
            ptable = doc.add_table(rows=1, cols=3)
            ptable.style = "Table Grid"
            hdr = ptable.rows[0].cells
            for i, name in enumerate(["词对", "共现文本数", "PMI"]):
                hdr[i].text = name
            for r in pair_rows:
                cells = ptable.add_row().cells
                cells[0].text = f"{r['source']} — {r['target']}"
                cells[1].text = str(r["count"])
                cells[2].text = f"{r['pmi']:.2f}"
    for which, caption in (
        ("positive", "正面讨论词云"),
        ("negative", "负面讨论词云"),
        ("worst_dim", "负面率最高维度词云"),
    ):
        wc_bytes = wordcloud_png_bytes(s, which)
        if wc_bytes:
            doc.add_picture(BytesIO(wc_bytes), width=Inches(6.2))
            run = doc.add_paragraph().add_run(caption)
            run.bold = True
    wc_insight = bundle.chart_insights.get("wordcloud", "")
    if wc_insight:
        p = doc.add_paragraph(wc_insight)
        p.paragraph_format.first_line_indent = Pt(24)

    # F-010（2026-08-26）：解读与建议段（与 HTML/Excel 同源：structured_summary）
    if bundle.structured_summary:
        _ss = bundle.structured_summary
        doc.add_heading("六、解读与建议", level=1)
        doc.add_paragraph(f"一句话结论：{_ss.get('overall', '')}")
        _pos = _ss.get("positive") or {}
        _neg = _ss.get("negative") or {}
        if _pos:
            p = doc.add_paragraph()
            run = p.add_run(f"正面反馈：占比 {_pos.get('ratio', 0) * 100:.1f}%")
            run.bold = True
            for _ph in _pos.get("phrases") or []:
                doc.add_paragraph(
                    f"  - {_ph.get('phrase', '')}"
                    f"（{_ph.get('count', 0)} 条/{_ph.get('ratio', 0) * 100:.0f}%）"
                )
        if _neg:
            p = doc.add_paragraph()
            run = p.add_run(f"负面反馈：占比 {_neg.get('ratio', 0) * 100:.1f}%")
            run.bold = True
            for _ph in _neg.get("phrases") or []:
                doc.add_paragraph(
                    f"  - {_ph.get('phrase', '')}"
                    f"（{_ph.get('count', 0)} 条/{_ph.get('ratio', 0) * 100:.0f}%）"
                )
        for _issue in _ss.get("top_issues") or []:
            p = doc.add_paragraph()
            run = p.add_run(
                f"重点问题：{_issue.get('name', '')}"
                f"（负面率 {_issue.get('rate', 0) * 100:.0f}%，n={_issue.get('count', 0)}）"
            )
            run.bold = True
            if _issue.get("cause"):
                doc.add_paragraph(f"可能原因（规则推测）：{_issue['cause']}")
            if _issue.get("direction"):
                doc.add_paragraph(f"建议：{_issue['direction']}")
        if _ss.get("improvements"):
            p = doc.add_paragraph()
            run = p.add_run("改进建议")
            run.bold = True
            for _im in _ss["improvements"]:
                doc.add_paragraph(f"- {_im}")
        doc.add_paragraph(
            "注：原因基于统计特征的规则推测，非 AI 归因，请结合报告原文验证。"
        )

    doc.add_heading(f"七、{findings_section_title(bundle.insight_mode)}", level=1)
    if bundle.findings:
        for f in bundle.findings:
            p = doc.add_paragraph()
            run = p.add_run(
                f"{display_finding_id(f.get('id', ''))} {f.get('claim', '')}"
            )
            run.bold = True
            for rid in (f.get("evidence_refs") or []):
                c = next((e for e in bundle.evidence if e.get("id") == rid), None)
                if not c:
                    continue
                if c.get("kind") == "stat":
                    doc.add_paragraph(
                        f"▸ 统计：{c['text']}（来源：报告统计 · n={c.get('n')}）"
                    )
                    continue
                judge = "LLM 判定" if c["judge"] == "llm" else "词典判定 · 仅供参考"
                label = judge
                if c.get("need_review"):
                    label += " · 待复核"
                doc.add_paragraph(
                    f"▸ 原文：「{c['text']}」"
                    f"（{c['platform']} · {c.get('date') or '日期未知'} · {label}）"
                )
            if f.get("action"):
                doc.add_paragraph(f"建议：{display_action(f.get('action'))}")
    elif bundle.conclusion:
        for line in bundle.conclusion.split("\n"):
            if line.strip():
                doc.add_paragraph(line)
    else:
        for line in bundle.report_text.split("\n"):
            p = doc.add_paragraph(line)
            p.paragraph_format.first_line_indent = Pt(24)

    doc.add_heading("八、概览", level=1)
    for line in bundle.report_text.split("\n"):
        p = doc.add_paragraph(line)
        p.paragraph_format.first_line_indent = Pt(24)

    if bundle.warnings:
        doc.add_heading("九、注意事项", level=1)
        for w in bundle.warnings:
            doc.add_paragraph(f"- {w}")

    _append_keyword_appendix(doc, bundle)

    out = BytesIO()
    doc.save(out)
    out.seek(0)
    return out
