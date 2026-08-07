# -*- coding: utf-8 -*-
"""人工抽检校准应用：把人工复核结果写入黄金集 v1.1。

口径（2026-08-07 定版，见 docs/抽样与标注规范.md §十一）：
    - 人工填写的字段（整条情感 / 强度 / 维度情感 / 是否相关）一律以人工为准
      （黄金集 = 人工裁判集）
    - 维度情感列：填写 JSON 或 "{}" 视为明确值并覆盖；留空则保留原 AI 标注
    - 变更写入 finalize_note（追加"人工校准"记录），可追溯

处理顺序：finalize_golden_set.py 定稿 → 本脚本人工校准 → benchmark_golden.py 评测。

用法：
    python tests/apply_human_review.py

输出：
    data/datasets/golden_set_v1.csv         工作版（含来源信息）
    tests/fixtures/golden_set_v1.csv        去 URL 入库版（跨机器复现）
"""

from __future__ import annotations

import csv
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

if sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "data" / "datasets" / "golden_set_v1.csv"
FIXTURE = ROOT / "tests" / "fixtures" / "golden_set_v1.csv"
XLSX = ROOT / "data" / "datasets" / "human_review_worksheet_v1.xlsx"

# 抽检表列位（按 build 时固定顺序，避免中文列名在控制台编码下抖动）
COL = {
    "text_id": 1,
    "ai_sentiment": 6,
    "ai_intensity": 7,
    "ai_dim": 8,
    "ai_relevant": 10,
    "human_sentiment": 11,
    "human_intensity": 12,
    "human_dim": 13,
    "human_relevant": 14,
    "agree_flag": 15,
    "remark": 16,
}


def norm(v) -> str:
    return "" if v is None else str(v).strip()


def parse_dim(raw: str) -> dict | None:
    """解析维度情感 JSON；解析失败返回 None。"""
    raw = raw.strip()
    if not raw:
        return None
    try:
        data = json.loads(raw)
    except Exception:
        return None
    return data if isinstance(data, dict) else None


def main() -> None:
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    xl_rows = [r for r in ws.iter_rows(values_only=True) if r[COL["text_id"]]]
    wb.close()

    review: dict[str, dict] = {}
    for r in xl_rows[1:]:
        review[str(r[COL["text_id"]])] = {
            "human_sentiment": norm(r[COL["human_sentiment"]]).lower(),
            "human_intensity": norm(r[COL["human_intensity"]]),
            "human_dim": norm(r[COL["human_dim"]]),
            "human_relevant": norm(r[COL["human_relevant"]]).lower(),
        }

    with open(GOLDEN, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    cols = list(rows[0].keys())

    changed: list[dict] = []
    dim_exact = dim_n = 0
    for row in rows:
        rev = review.get(row["text_id"])
        if not rev:
            continue
        notes: list[str] = []

        old_sent = row["sentiment"]
        if rev["human_sentiment"] and rev["human_sentiment"] != old_sent:
            row["sentiment"] = rev["human_sentiment"]
            notes.append(f"情感 {old_sent}->{rev['human_sentiment']}")
        if rev["human_intensity"]:
            old_i = row["intensity"]
            if rev["human_intensity"] != old_i:
                row["intensity"] = rev["human_intensity"]
                notes.append(f"强度 {old_i}->{rev['human_intensity']}")
        if rev["human_relevant"]:
            old_rel = row["relevant"]
            if rev["human_relevant"] != old_rel:
                row["relevant"] = rev["human_relevant"]
                notes.append(f"相关 {old_rel}->{rev['human_relevant']}")

        # 维度情感：填写（含 {}）即覆盖；留空保留原标注
        if rev["human_dim"]:
            human_dims = parse_dim(rev["human_dim"])
            old_dims = row["dimension_sentiments"]
            if human_dims is None:
                notes.append(f"维度解析失败，保留原值（{rev['human_dim'][:40]}）")
            else:
                row["dimension_sentiments"] = json.dumps(human_dims, ensure_ascii=False)
                if human_dims:
                    dim_n += 1
                    try:
                        if set(human_dims) == set(json.loads(old_dims or "{}")):
                            dim_exact += 1
                    except Exception:
                        pass
                if human_dims != json.loads(old_dims or "{}"):
                    notes.append(f"维度 {old_dims or '{}'}->{row['dimension_sentiments']}")

        if notes:
            old_note = row.get("finalize_note") or ""
            row["finalize_note"] = (old_note + "；" if old_note else "") + "人工校准:" + "，".join(notes)
            changed.append({"text_id": row["text_id"], "notes": notes})

    with open(GOLDEN, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)

    fixture_cols = [c for c in cols if c not in ("url", "likes", "time")]
    FIXTURE.parent.mkdir(parents=True, exist_ok=True)
    with open(FIXTURE, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fixture_cols)
        w.writeheader()
        for row in rows:
            w.writerow({c: row[c] for c in fixture_cols})

    print(f"黄金集 v1.1 已写入：{len(rows)} 条（{len(changed)} 条被人工修正）")
    print("相关分布:", dict(Counter(r["relevant"] for r in rows)))
    print("情感分布:", dict(Counter(r["sentiment"] for r in rows)))
    print("维度级一致（抽检中已填写维度样本）:", f"{dim_exact}/{dim_n}")
    for c in changed:
        print(" ", c["text_id"], "->", c["notes"])


if __name__ == "__main__":
    main()
