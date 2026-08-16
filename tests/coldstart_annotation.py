# -*- coding: utf-8 -*-
"""2.5 冷启动裁判工具链（领域无关）。

用法：
    python tests/coldstart_annotation.py --worksheet --domain digital3c \
        --out data/datasets/coldstart_digital3c_worksheet.xlsx
    python tests/coldstart_annotation.py --demo-skeleton --domain digital3c \
        --source data/archive/online_shopping_10_cats.csv --n 60 \
        --out data/datasets/coldstart_digital3c_golden.csv
    python tests/coldstart_annotation.py --finalize --sheet <xlsx> \
        --out data/datasets/coldstart_digital3c_golden.csv --domain digital3c
    python tests/coldstart_annotation.py --from-report data/reports/<任务目录> \
        [data/reports/<任务目录2> ...] --domain digital3c --dims-json <提案.json> \
        --n 120 --out <标注表.xlsx>
    python tests/coldstart_annotation.py --compare --primary <主标.xlsx> \
        --secondary <副标.xlsx> --domain digital3c
    python tests/coldstart_annotation.py --finalize --primary <主标.xlsx> \
        [--secondary <副标.xlsx>] --out <golden.csv> --domain digital3c

流程：空白标注表（维度列按领域 schema 动态渲染）→ 双 AI/人工标注 →
--finalize 定版 golden → benchmark --golden <csv> --mode-key domain_<id>_lexicon
→ 新领域独立基线。
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.domains.loader import load_domain  # noqa: E402

BASE_COLUMNS = [
    "text_id", "原文", "平台", "品牌/主题", "关键词", "文本类型",
    "情感(整条)", "强度(1-5)", "语言现象", "是否相关", "备注", "标注人", "标注日期",
]
GOLD_REV_COL = "是否参考过模型判定"
# V1 hold-out 纪律（2.6 验证纪律）：每份卷子最多验收次数；冻结后只读
HOLD_OUT_USES_MAX = 2
STRATUM_TOLERANCE_PP = 10.0


def _schema_dims(domain_id: str, dims_json: Path | None = None) -> list:
    """维度来源：--dims-json 提案维度（骨架/标注表，正式 schema 未确认前可用）或已落库 schema。"""
    if dims_json and dims_json.exists():
        proposal = json.loads(dims_json.read_text(encoding="utf-8"))
        return [type("D", (), {"id": d["id"], "name": d["name"]}) for d in proposal["dimensions"]]
    return load_domain(domain_id).dimensions


def gen_worksheet(domain_id: str, out: Path, samples: list[dict] | None = None,
                  dims_json: Path | None = None) -> None:
    dims = _schema_dims(domain_id, dims_json)
    headers = BASE_COLUMNS[:6] + [d.name for d in dims] + BASE_COLUMNS[6:]
    wb = Workbook()
    ws = wb.active
    ws.title = "标注"
    ws.append(headers)
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="E2EFDA")
        c.font = Font(bold=True)
    for s in samples or []:
        row = [s.get(h, "") for h in headers]
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    for i, wd in enumerate([14, 60, 10, 16, 16, 10] + [14] * len(dims) + [18, 10, 20, 10, 10], 1):
        ws.column_dimensions[chr(64 + i)].width = wd
    guide = wb.create_sheet("填写说明")
    for line in [
        f"新领域「{domain_id}」冷启动标注表：维度列按 schema 动态生成（{len(dims)} 个）。",
        "情感(整条)：positive/negative/neutral；强度 1~5（neutral=1）。",
        "维度列：只填 positive/negative；无明确褒贬留空（对齐标注规范 §五）。",
        "是否相关：yes/no；no 的样本情感标 neutral、维度留空（不计入主评分）。",
        "定版：python tests/coldstart_annotation.py --finalize --sheet <本文件> --out <golden.csv>",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    wb.save(out)
    print(f"标注表已生成：{out}（{len(dims)} 个动态维度列）")


def finalize(sheet: Path, out: Path, domain_id: str, dims_json: Path | None = None) -> None:
    dims = _schema_dims(domain_id, dims_json)
    wb = load_workbook(sheet, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    headers = [str(h) if h else "" for h in next(it)]
    has_gold_rev = GOLD_REV_COL in headers
    golden = []
    for i, rec in enumerate(it, 1):
        if not rec or not str(rec[headers.index("原文")] or "").strip():
            continue
        sent = str(rec[headers.index("情感(整条)")] or "").strip().lower()
        if not sent:
            continue
        dims_out = {}
        for d in dims:
            v = str(rec[headers.index(d.name)] or "").strip().lower()
            if v in ("positive", "negative"):
                dims_out[d.name] = v
        relevant = str(rec[headers.index("是否相关")] or "").strip().lower() or "yes"
        golden.append({
            "text_id": str(rec[headers.index("text_id")] or f"{domain_id}_C{i:04d}"),
            "text": str(rec[headers.index("原文")] or ""),
            "platform": str(rec[headers.index("平台")] or "").strip() or "demo",
            "domain": domain_id,
            "brand": str(rec[headers.index("品牌/主题")] or "").strip(),
            "keyword": str(rec[headers.index("关键词")] or "").strip(),
            "kind": str(rec[headers.index("文本类型")] or "").strip() or "评论",
            "sentiment": sent,
            "intensity": str(rec[headers.index("强度(1-5)")] or "").strip() or ("1" if sent == "neutral" else "3"),
            "dimension_sentiments": json.dumps(dims_out, ensure_ascii=False),
            "language_flags": str(rec[headers.index("语言现象")] or "").strip(),
            "relevant": relevant,
            "remark": str(rec[headers.index("备注")] or "").strip(),
            "annotator": str(rec[headers.index("标注人")] or "").strip(),
            "annotated_at": str(rec[headers.index("标注日期")] or "").strip(),
            "gold_revised_by_model": (
                str(rec[headers.index(GOLD_REV_COL)] or "").strip()
                if has_gold_rev else ""
            ),
        })
    wb.close()
    if not golden:
        raise ValueError("标注表没有已填情感的行，无法定版")
    fields = list(golden[0].keys())
    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(golden)
    print(f"定版 golden：{out}（{len(golden)} 条，领域 {domain_id}）")


def demo_skeleton(domain_id: str, source: Path, n: int, out: Path,
                  worksheet_out: Path | None = None, dims_json: Path | None = None) -> None:
    """从公开语料抽数码类骨架（情感来自原标签，维度留空，待人工/双 AI 标注）。"""
    rows = list(csv.DictReader(open(source, encoding="utf-8-sig")))
    digital = ["手机", "计算机", "平板"]
    pool = [r for r in rows if r.get("cat") in digital]
    random.Random(42).shuffle(pool)
    picked = pool[:n]
    samples = []
    for i, r in enumerate(picked, 1):
        label = int(r.get("label") or 0)
        samples.append({
            "text_id": f"{domain_id}_C{i:04d}",
            "原文": r.get("review", "").strip(),
            "平台": "websearch",
            "品牌/主题": "",
            "关键词": r.get("cat", ""),
            "文本类型": "评论",
        })
    gen_worksheet(domain_id, worksheet_out or out.with_name(out.stem + "_worksheet.xlsx"),
                  samples, dims_json)
    golden = []
    for i, s in enumerate(samples, 1):
        label = int(picked[i - 1].get("label") or 0)
        sent = "positive" if label == 1 else "negative"
        golden.append({
            "text_id": s["text_id"], "text": s["原文"], "platform": "websearch",
            "domain": domain_id, "brand": "", "keyword": s["关键词"], "kind": "评论",
            "sentiment": sent, "intensity": "3",
            "dimension_sentiments": "{}", "language_flags": "",
            "relevant": "yes", "remark": "骨架样本（公开语料，维度待标注）",
            "annotator": "", "annotated_at": "",
        })
    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(golden[0].keys()))
        w.writeheader()
        w.writerows(golden)
    print(f"骨架 golden：{out}（{len(golden)} 条，{picked[0]['cat'] if picked else ''} 等数码类）")


def extract_report_items(report_dir: Path, domain_id: str) -> list[dict]:
    """从真实任务报告提取目标领域样本（含 pub_date，供 hold-out 时间窗分层）。"""
    items = []
    for p in sorted(report_dir.glob("**/result.json")):
        bundle = json.loads(p.read_text(encoding="utf-8"))
        subject = (bundle.get("plan") or {}).get("subject", "")
        if (bundle.get("plan") or {}).get("domain_id") != domain_id:
            continue  # 只采样目标领域的任务（目录可能混有其他领域报告）
        for it in bundle.get("coded_items", []):
            text = (it.get("text") or "").strip()
            if not text:
                continue
            items.append({
                "text_id": it.get("text_id", ""),
                "原文": text,
                "平台": it.get("platform", ""),
                "品牌/主题": subject,
                "关键词": it.get("keyword", ""),
                "文本类型": "评论" if ":comment" in it.get("text_id", "") else "帖子正文",
                "pub_date": it.get("pub_date", ""),
            })
    seen, uniq = set(), []
    for s in items:
        if s["原文"] in seen:
            continue
        seen.add(s["原文"])
        uniq.append(s)
    return uniq


def extract_report_items_many(report_dirs: list[Path], domain_id: str) -> list[dict]:
    """跨任务聚合提取（养肥/难例回流多目录用）：按原文跨任务去重。"""
    seen: set[str] = set()
    pooled: list[dict] = []
    for rd in report_dirs:
        for s in extract_report_items(rd, domain_id):
            text = s.get("原文") or ""
            if text in seen:
                continue
            seen.add(text)
            pooled.append(s)
    return pooled


def load_excluded_texts(paths: list[Path]) -> set[str]:
    """读取 golden/hold-out CSV 的原文集合（--exclude-golden 用）。

    新卷切分必须排除已入裁判集的文本，否则"新样本"混入旧样本，
    违反 hold-out 同分布/独立原则。
    """
    out: set[str] = set()
    for p in paths:
        if not p.exists():
            continue
        for r in csv.DictReader(open(p, encoding='utf-8-sig')):
            t = (r.get("text") or "").strip()
            if t:
                out.add(t)
    return out


def from_report(report_dirs: list[Path], domain_id: str, n: int, out: Path,
                dims_json: Path | None = None,
                exclude_golden: list[Path] | None = None) -> None:
    """从真实任务报告（可多目录）采样生成独立标注表（情感/维度列留空）。"""
    uniq = extract_report_items_many(report_dirs, domain_id)
    if exclude_golden:
        excluded = load_excluded_texts(exclude_golden)
        before = len(uniq)
        uniq = [s for s in uniq if (s.get("原文") or "").strip() not in excluded]
        if len(uniq) < before:
            print(f"[排除] 剔除与 golden/hold-out 重叠文本 {before - len(uniq)} 条")
    random.Random(42).shuffle(uniq)
    samples = uniq[:n]
    if len(samples) < n:
        print(f"[提示] 报告内去重后仅 {len(samples)} 条（目标 {n}），可多品牌/多渠道补采")
    gen_worksheet(domain_id, out, samples, dims_json)
    print(f"已从报告采样 {len(samples)} 条 → 独立标注表：{out}")


def _pub_date(v: str) -> date | None:
    try:
        return date.fromisoformat((v or "").strip())
    except ValueError:
        return None


def _time_bucket(pub_date: str, min_d: date | None, max_d: date | None) -> str:
    """时间窗三分桶：把报告时间范围均分 3 段，按 pub_date 落桶（'0'/'1'/'2'）；
    缺失/非法 → 'unknown'。统一字符串便于排序与配额校验。"""
    d = _pub_date(pub_date)
    if d is None or min_d is None or max_d is None:
        return "unknown"
    span = (max_d - min_d).days
    if span <= 0:
        return "0"
    return str(min(2, (d - min_d).days * 3 // span))


def split_holdout(items: list[dict], n: int, seed: int = 42) -> tuple[list[dict], list[dict], dict]:
    """hold-out 分层切分：同分布约束 = (平台, 品牌) 配额 + 时间窗三分桶。

    返回 (holdout, iteration, 校验报告)。配额偏差 >10pp 或时间桶缺失（池内该桶 ≥10 条）
    视为不合格并抛错——否则 hold-out 提升可能只是分布差异，不是模型变好。
    """
    if n <= 0:
        raise ValueError("hold-out 数量必须 > 0（建议 30~50 条）")
    if not items:
        raise ValueError("报告内没有可切分样本")
    dates = [d for d in (_pub_date(i.get("pub_date")) for i in items) if d]
    min_d, max_d = (min(dates), max(dates)) if dates else (None, None)
    strata: dict[tuple, list[dict]] = defaultdict(list)
    for it in items:
        key = (
            it.get("平台") or "unknown",
            it.get("品牌/主题") or "unknown",
            _time_bucket(it.get("pub_date"), min_d, max_d),
        )
        strata[key].append(it)
    total = len(items)
    n = min(n, total)
    # 目标配额：按分层占比（最大余数法）
    targets = {k: len(v) / total * n for k, v in strata.items()}
    pick = {k: int(t) for k, t in targets.items()}
    rem = n - sum(pick.values())
    for k in sorted(targets, key=lambda x: -targets[x] % 1)[:max(rem, 0)]:
        pick[k] += 1
    rng = random.Random(seed)
    holdout, iteration, shortfall = [], [], {}
    for key, vals in strata.items():
        rng.shuffle(vals)
        k = min(pick[key], len(vals))
        if k < pick[key]:
            shortfall[str(key)] = (pick[key], len(vals))
        holdout.extend(vals[:k])
        iteration.extend(vals[k:])
    # 配额校验
    def share(rows, key):
        s = sum(1 for r in rows if (r.get("平台") or "unknown",
                                    r.get("品牌/主题") or "unknown",
                                    _time_bucket(r.get("pub_date"), min_d, max_d)) == key)
        return s / len(rows) * 100 if rows else 0.0
    issues = []
    for key, vals in strata.items():
        pool_share = len(vals) / total * 100
        hold_share = share(holdout, key)
        if abs(hold_share - pool_share) > STRATUM_TOLERANCE_PP:
            issues.append(f"分层 {key} 配额偏差 {hold_share - pool_share:+.1f}pp"
                          f"（池 {pool_share:.1f}% vs hold-out {hold_share:.1f}%）")
    buckets = sorted({str(k[2]) for k in strata})
    for b in buckets:
        pool_n = sum(len(v) for k, v in strata.items() if k[2] == b)
        hold_n = sum(1 for r in holdout
                     if _time_bucket(r.get("pub_date"), min_d, max_d) == b)
        if pool_n >= 10 and hold_n == 0:
            issues.append(f"时间桶 {b} 缺失（池 {pool_n} 条，hold-out 0 条）")
    report = {
        "seed": seed,
        "n_target": n,
        "n_pool": total,
        "n_holdout": len(holdout),
        "n_iteration": len(iteration),
        "time_range": [min_d.isoformat() if min_d else None,
                       max_d.isoformat() if max_d else None],
        "strata": {str(k): {"pool": len(v), "holdout": share(holdout, k)}
                   for k, v in strata.items()},
        "shortfall": shortfall,
        "issues": issues,
        "note": "同分布约束 = (平台, 品牌) 配额（容差 10pp）+ 时间窗三分桶覆盖",
    }
    if issues:
        raise ValueError("hold-out 切分未通过同分布校验：\n  " + "\n  ".join(issues))
    return holdout, iteration, report


def _holdout_marker(golden: Path) -> Path:
    return golden.with_suffix(golden.suffix + ".frozen.json")


def freeze_holdout(golden: Path, force: bool = False) -> Path:
    """hold-out 定版后冻结：写指纹 + 冻结标记，只读（重切需 --force 作废旧卷）。"""
    if not golden.exists():
        raise ValueError(f"hold-out golden 不存在：{golden}")
    marker = _holdout_marker(golden)
    if marker.exists() and not force:
        raise ValueError(f"{marker.name} 已存在（hold-out 冻结后只读）；"
                         "确需重切请 --force（旧卷作废，uses 清零）")
    from app.core import eval_store
    marker.write_text(json.dumps({
        "frozen": True,
        "fingerprint": eval_store.golden_fingerprint(golden),
        "frozen_at": datetime.now().isoformat(timespec="seconds"),
        "uses": 0,
        "max_uses": HOLD_OUT_USES_MAX,
        "note": "冻结后只读：hold-out 错误不进失败模式、不据此改规则；"
                "每里程碑重切、≤2 次验收（benchmark --golden holdout_*.csv 自动计数）",
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    return marker


def record_holdout_use(golden: Path) -> dict:
    """记录一次 hold-out 验收使用；达到上限（默认 2 次）拒绝继续（需重切新卷）。"""
    marker = _holdout_marker(golden)
    if not marker.exists():
        return {"ok": False, "error": "hold-out 未冻结（先运行 --freeze-holdout）"}
    try:
        data = json.loads(marker.read_text(encoding="utf-8"))
    except ValueError:
        return {"ok": False, "error": f"冻结标记损坏：{marker.name}"}
    if not data.get("frozen"):
        return {"ok": False, "error": "冻结标记非冻结态"}
    # 冻结只读的运行时校验：内容指纹变化说明卷子被改过，作废
    from app.core import eval_store

    cur_fp = eval_store.golden_fingerprint(golden)
    fp = data.get("fingerprint")
    if fp and cur_fp != fp:
        return {"ok": False, "blocked": True, "error":
                f"hold-out 内容已变更（冻结指纹 {fp} ≠ 当前 {cur_fp}）；"
                "该卷作废，需 --freeze-holdout --force 重切新卷"}
    max_uses = int(data.get("max_uses") or HOLD_OUT_USES_MAX)
    used = int(data.get("uses") or 0)
    if used >= max_uses:
        return {"ok": False, "blocked": True, "uses": used, "max_uses": max_uses,
                "error": f"hold-out 已验收 {used}/{max_uses} 次，达上限；需重切新卷"}
    data["uses"] = used + 1
    data["last_used_at"] = datetime.now().isoformat(timespec="seconds")
    marker.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"ok": True, "uses": data["uses"], "max_uses": max_uses}


def load_annotated(path: Path, domain_id: str) -> list[dict]:
    dims = _schema_dims(domain_id)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h) if h else "" for h in next(it)]
    has_gold_rev = GOLD_REV_COL in hdr
    out = []
    for rec in it:
        if not rec or not str(rec[hdr.index("原文")] or "").strip():
            continue
        dim_map = {}
        for d in dims:
            v = str(rec[hdr.index(d.name)] or "").strip().lower()
            if v in ("positive", "negative"):
                dim_map[d.name] = v
        out.append({
            "text_id": str(rec[hdr.index("text_id")] or "").strip(),
            "text": str(rec[hdr.index("原文")] or "").strip(),
            "platform": str(rec[hdr.index("平台")] or "").strip(),
            "brand": str(rec[hdr.index("品牌/主题")] or "").strip(),
            "keyword": str(rec[hdr.index("关键词")] or "").strip(),
            "kind": str(rec[hdr.index("文本类型")] or "").strip(),
            "sentiment": str(rec[hdr.index("情感(整条)")] or "").strip().lower(),
            "intensity": str(rec[hdr.index("强度(1-5)")] or "").strip(),
            "dims": dim_map,
            "flags": str(rec[hdr.index("语言现象")] or "").strip(),
            "relevant": str(rec[hdr.index("是否相关")] or "").strip().lower(),
            "remark": str(rec[hdr.index("备注")] or "").strip(),
            "annotator": str(rec[hdr.index("标注人")] or "").strip(),
            "annotated_at": str(rec[hdr.index("标注日期")] or "").strip(),
            "gold_revised_by_model": (
                str(rec[hdr.index(GOLD_REV_COL)] or "").strip()
                if has_gold_rev else ""
            ),
        })
    wb.close()
    return out


def compare_annotations(primary: list[dict], secondary: list[dict]) -> tuple[dict, list[str]]:
    from tests.edge_annotation import cohen_kappa

    p = {(r["text_id"], r["text"]): r for r in primary}
    s = {(r["text_id"], r["text"]): r for r in secondary}
    uids = [k for k in p if k in s]
    if not uids:
        raise ValueError("主标/副标没有交集")
    sp = [p[t]["sentiment"] for t in uids]
    ss = [s[t]["sentiment"] for t in uids]
    rp = [p[t]["relevant"] for t in uids]
    rs = [s[t]["relevant"] for t in uids]
    dim_n = dim_agree = 0
    for t in uids:
        dp, ds = p[t]["dims"], s[t]["dims"]
        if dp or ds:
            dim_n += 1
            dim_agree += dp == ds
    disputed = [t[0] for t in uids
                if p[t]["sentiment"] != s[t]["sentiment"]
                or p[t]["relevant"] != s[t]["relevant"]]
    return {
        "n": len(uids),
        "sentiment_agreement": round(sum(a == b for a, b in zip(sp, ss)) / len(uids), 4),
        "kappa": cohen_kappa(sp, ss),
        "relevance_agreement": round(sum(a == b for a, b in zip(rp, rs)) / len(uids), 4),
        "dimension_agreement": round(dim_agree / dim_n, 4) if dim_n else None,
        "dimension_n": dim_n,
        "disputed_n": len(disputed),
        "disputed_ids": disputed,
        "note": "按 (text_id, 原文) 复合键比对（同帖多条评论共用 text_id 的历史数据模型）",
    }, disputed


def load_human_decisions(path: Path) -> dict[str, dict]:
    """解析人工复核表（分歧/高歧义）：→ {sentiment, relevant}。

    键：text_id 唯一时用 text_id；同帖多评论共用 text_id 时用 (text_id, 原文)
    复合键（避免把一条评论的裁决误套到同帖其他评论）。列契约：
    「text_id」「原文」「人工判定情感」「人工判定相关」；留空 = 不覆盖。
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    out: dict = {}
    try:
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        hdr = [str(h) if h else "" for h in next(it, [])]
        has_sent = "人工判定情感" in hdr
        has_rel = "人工判定相关" in hdr
        recs: list[tuple[str, str, str, str]] = []
        from collections import Counter

        tid_counts: Counter = Counter()
        for rec in it:
            if not rec or not str(rec[0] if rec else "").strip():
                continue
            tid = str(rec[hdr.index("text_id")]).strip()
            if not tid:
                continue
            sent = str(rec[hdr.index("人工判定情感")] or "").strip().lower() if has_sent else ""
            rel = str(rec[hdr.index("人工判定相关")] or "").strip().lower() if has_rel else ""
            if sent or rel:
                text = str(rec[hdr.index("原文")] or "").strip() if "原文" in hdr else ""
                recs.append((tid, text, sent, rel))
                tid_counts[tid] += 1
        for tid, text, sent, rel in recs:
            key = tid if tid_counts[tid] == 1 else (tid, text)
            out[key] = {"sentiment": sent, "relevant": rel}
    finally:
        wb.close()
    return out


def gen_dispute_review(primary: list[dict], secondary: list[dict], out: Path,
                       domain_id: str) -> int:
    """生成分歧/高歧义人工复核表（情感/相关不一致行，主标 vs 副标对照）。

    维度分歧按 2.5 决策以主标为准，不进入复核表。返回分歧行数。
    """
    s_map = {(r["text_id"], r["text"]): r for r in secondary}
    rows: list[tuple[dict, dict]] = []
    for a in primary:
        b = s_map.get((a["text_id"], a["text"]))
        if b and (a["sentiment"] != b["sentiment"] or a["relevant"] != b["relevant"]):
            rows.append((a, b))
    cols = ["序号", "text_id", "原文", "平台", "品牌/主题",
            "主标情感", "副标情感", "主标相关", "副标相关",
            "主标维度", "副标维度", "人工判定情感", "人工判定相关", "备注"]
    wb = Workbook()
    ws = wb.active
    ws.title = "分歧复核"
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="E2EFDA")
        cell.font = Font(bold=True)
    for i, (a, b) in enumerate(rows, 1):
        ws.append([
            i, a["text_id"], a["text"], a["platform"], a["brand"],
            a["sentiment"] or "", b["sentiment"] or "",
            a["relevant"] or "", b["relevant"] or "",
            ",".join(f"{k}:{v}" for k, v in a["dims"].items()),
            ",".join(f"{k}:{v}" for k, v in b["dims"].items()),
            "", "", "",
        ])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    for i, w in enumerate([6, 14, 55, 10, 14, 10, 10, 8, 8, 26, 26, 12, 10, 20], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    guide = wb.create_sheet("填写说明")
    for line in [
        f"领域 {domain_id} 分歧/高歧义人工复核表（情感/相关不一致，共 {len(rows)} 条）。",
        "一、留空 = 认可主标（默认）；不同意主标请在「人工判定情感/相关」填最终值。",
        "二、维度分歧按 2.5 决策以主标为准，本表不重复核对维度。",
        "三、填完发回，或回复「全按主标」；定版：coldstart --finalize --human-review <本表>",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"复核表已生成：{out}（{len(rows)} 条）")
    return len(rows)


def finalize_merged(primary_path: Path, secondary_path: Path | None, out: Path,
                    domain_id: str, report_out: Path | None = None,
                    exclude_disputed: bool = False,
                    fixture_out: Path | None = None,
                    human_review: Path | None = None) -> None:
    """主标定版：主标为准 + 副标分歧记录；规范归一（relevant=no 清维度、neutral 强度=1）。"""
    primary = load_annotated(primary_path, domain_id)
    secondary = load_annotated(secondary_path, domain_id) if secondary_path else []
    human = load_human_decisions(human_review) if human_review else {}
    s = {(r["text_id"], r["text"]): r for r in secondary}
    compare_report, _ = compare_annotations(primary, secondary) if secondary else ({}, [])
    golden, notes = [], []
    for r in primary:
        tid = r["text_id"]
        uid = (tid, r["text"])
        sent, relevant = r["sentiment"], r["relevant"] or "yes"
        dims = dict(r["dims"])
        note = []
        h = human.get((tid, r["text"])) or human.get(tid)
        if h:
            if h["sentiment"] and h["sentiment"] != sent:
                note.append(f"人工复核情感:{sent or '空'}→{h['sentiment']}")
                sent = h["sentiment"]
            if h["relevant"] and h["relevant"] != relevant:
                note.append(f"人工复核相关:{relevant or '空'}→{h['relevant']}")
                relevant = h["relevant"]
        if uid in s and s[uid]["sentiment"] != sent:
            note.append(f"副标分歧:{s[uid]['sentiment'] or '空'}")
        if relevant == "no":
            if dims:
                note.append("相关=no 维度清空（规范 §五）")
                dims = {}
            if sent != "neutral":
                note.append("相关=no 情感归一 neutral")
                sent = "neutral"
        if sent == "neutral" and r["intensity"] not in ("", "1"):
            note.append("neutral 强度归一 1")
        if note:
            notes.append(tid)
        golden.append({
            "text_id": tid, "text": r["text"], "platform": r["platform"] or "demo",
            "domain": domain_id, "brand": r["brand"], "keyword": r["keyword"],
            "kind": r["kind"] or "评论",
            "sentiment": sent,
            "intensity": "1" if sent == "neutral" else (r["intensity"] or "3"),
            "dimension_sentiments": json.dumps(dims, ensure_ascii=False),
            "language_flags": r["flags"], "relevant": relevant,
            "remark": ("；".join(note) + ("；" if note else "") + r["remark"]) if note else r["remark"],
            "annotator": r["annotator"], "annotated_at": r["annotated_at"],
            "gold_revised_by_model": r.get("gold_revised_by_model", ""),
        })
    if exclude_disputed and secondary:
        disputed = set(compare_report.get("disputed_ids") or [])
        keep = [g for g in golden if g["text_id"] not in disputed]
        print(f"[提示] --exclude-disputed：剔除 {len(golden) - len(keep)} 条分歧行")
        golden = keep
    if not golden:
        raise ValueError("定版后无有效行")
    # 同帖多评论共用 text_id → 加后缀保证 golden 唯一（历史数据模型修正）
    seen: dict[str, int] = {}
    for g in golden:
        base_id = g["text_id"]
        seen[base_id] = seen.get(base_id, 0) + 1
        if seen[base_id] > 1:
            g["text_id"] = f"{base_id}#{seen[base_id]}"
    fields = list(golden[0].keys())
    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(golden)
    print(f"定版 golden：{out}（{len(golden)} 条，领域 {domain_id}；规范归一 {len(notes)} 条）")
    if fixture_out:
        drop_cols = {"url", "likes", "time"}
        f_cols = [c for c in fields if c not in drop_cols]
        fixture_out.parent.mkdir(parents=True, exist_ok=True)
        with open(fixture_out, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=f_cols)
            w.writeheader()
            for g in golden:
                w.writerow({c: g.get(c, "") for c in f_cols})
        print(f"入库版（去 URL）：{fixture_out}")
    if report_out:
        report_out.parent.mkdir(parents=True, exist_ok=True)
        report_out.write_text(
            json.dumps({"compare": compare_report, "normalized": notes},
                       ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"一致性/归一报告：{report_out}")


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="2.5 冷启动裁判工具链")
    ap.add_argument("--worksheet", action="store_true")
    ap.add_argument("--demo-skeleton", action="store_true")
    ap.add_argument("--from-report", type=Path, nargs="+")
    ap.add_argument("--compare", action="store_true")
    ap.add_argument("--review-sheet", action="store_true",
                    help="生成分歧/高歧义人工复核表（--primary/--secondary/--out）")
    ap.add_argument("--primary", type=Path)
    ap.add_argument("--secondary", type=Path)
    ap.add_argument("--report", type=Path)
    ap.add_argument("--exclude-disputed", action="store_true")
    ap.add_argument("--finalize", action="store_true")
    ap.add_argument("--holdout-split", action="store_true")
    ap.add_argument("--freeze-holdout", action="store_true")
    ap.add_argument("--record-use", action="store_true")
    ap.add_argument("--domain", required=True)
    ap.add_argument("--out", type=Path)
    ap.add_argument("--out-dir", type=Path)
    ap.add_argument("--golden", type=Path)
    ap.add_argument("--fixture-out", type=Path)
    ap.add_argument("--human-review", type=Path, help="人工复核表 xlsx（分歧/高歧义裁决）")
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--exclude-golden", type=Path, nargs="+",
                    help="切分/采样时排除的裁判集 CSV（原文匹配，防新卷混入旧样本）")
    ap.add_argument("--force", action="store_true")
    ap.add_argument("--source", type=Path)
    ap.add_argument("--sheet", type=Path)
    ap.add_argument("--dims-json", type=Path)
    ap.add_argument("--n", type=int, default=60)
    args = ap.parse_args()
    if args.worksheet:
        if not args.out:
            ap.error("--worksheet 需要 --out")
        gen_worksheet(args.domain, args.out, dims_json=args.dims_json)
        return 0
    if args.demo_skeleton:
        if not (args.out and args.source):
            ap.error("--demo-skeleton 需要 --out 与 --source")
        demo_skeleton(args.domain, args.source, args.n, args.out, dims_json=args.dims_json)
        return 0
    if args.holdout_split:
        if not (args.from_report and args.out_dir):
            ap.error("--holdout-split 需要 --from-report 与 --out-dir")
        items = extract_report_items_many(args.from_report, args.domain)
        if args.exclude_golden:
            excluded = load_excluded_texts(args.exclude_golden)
            before = len(items)
            items = [s for s in items if (s.get("原文") or "").strip() not in excluded]
            if len(items) < before:
                print(f"[排除] 剔除与 golden/hold-out 重叠文本 {before - len(items)} 条")
        holdout, iteration, rep = split_holdout(items, args.n, args.seed)
        args.out_dir.mkdir(parents=True, exist_ok=True)
        gen_worksheet(args.domain, args.out_dir / f"holdout_{args.domain}_worksheet.xlsx",
                      holdout, args.dims_json)
        gen_worksheet(args.domain, args.out_dir / f"iteration_{args.domain}_worksheet.xlsx",
                      iteration, args.dims_json)
        manifest = {
            "kind": "holdout_split",
            "domain": args.domain,
            "report_dir": str(args.from_report),
            "created_at": datetime.now().isoformat(timespec="seconds"),
            **rep,
        }
        mp = args.out_dir / f"holdout_split_{args.domain}_manifest.json"
        mp.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"hold-out {len(holdout)} 条 / 迭代集 {len(iteration)} 条 → {args.out_dir}")
        print(f"同分布校验：{('通过' if not rep['issues'] else '见 manifest issues')}")
        print(f"manifest：{mp}")
        print("后续：双 AI/人工标注 → --finalize 定版 holdout_<domain>.csv → "
              "--freeze-holdout 冻结 → benchmark --golden holdout_<domain>.csv")
        return 0
    if args.from_report:
        if not args.out:
            ap.error("--from-report 需要 --out")
        from_report(args.from_report, args.domain, args.n, args.out, args.dims_json,
                    args.exclude_golden)
        return 0
    if args.freeze_holdout:
        if not args.golden:
            ap.error("--freeze-holdout 需要 --golden")
        marker = freeze_holdout(args.golden, force=args.force)
        print(f"hold-out 已冻结只读：{marker}")
        return 0
    if args.record_use:
        if not args.golden:
            ap.error("--record-use 需要 --golden")
        res = record_holdout_use(args.golden)
        if res.get("ok"):
            print(f"hold-out 验收已记录：{res['uses']}/{res['max_uses']} 次")
            return 0
        print(f"hold-out 使用记录失败：{res.get('error')}")
        return 1
    if args.compare:
        if not (args.primary and args.secondary):
            ap.error("--compare 需要 --primary 与 --secondary")
        rep, _ = compare_annotations(
            load_annotated(args.primary, args.domain),
            load_annotated(args.secondary, args.domain))
        print("双 AI 一致性：")
        for k in ("n", "sentiment_agreement", "kappa", "relevance_agreement",
                  "dimension_agreement", "dimension_n", "disputed_n"):
            print(f"  {k}: {rep[k]}")
        return 0
    if args.review_sheet:
        if not (args.primary and args.secondary and args.out):
            ap.error("--review-sheet 需要 --primary 与 --secondary 与 --out")
        n = gen_dispute_review(
            load_annotated(args.primary, args.domain),
            load_annotated(args.secondary, args.domain),
            args.out, args.domain)
        print(f"分歧/高歧义 {n} 条 → 人工复核表已生成")
        return 0
    if args.finalize:
        if not args.out:
            ap.error("--finalize 需要 --out")
        if args.primary:
            finalize_merged(args.primary, args.secondary, args.out, args.domain,
                            args.report, args.exclude_disputed, args.fixture_out,
                            args.human_review)
        else:
            finalize(args.sheet, args.out, args.domain, args.dims_json)
        return 0
    ap.print_help()
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
