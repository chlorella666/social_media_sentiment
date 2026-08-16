# -*- coding: utf-8 -*-
"""2.5 领域扩展提案引擎测试：模板建议/证据解析/提案合并/确认表应用/统一加载/自定义评测键。"""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path
from unittest import mock

if sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.domains import proposer  # noqa: E402
from app.domains.loader import load_domain  # noqa: E402

EVIDENCE = ROOT / "app" / "domains" / "evidence" / "digital3c.md"


def test_suggest_template(tmp_path=None):
    tid, reasons, conf = proposer.suggest_template("数码3C", "手机相机电脑硬件")
    assert tid == "physical" and 0 < conf <= 1 and reasons
    assert proposer.suggest_template("某餐厅", "餐饮服务")[0] == "service"
    assert proposer.suggest_template("某游戏", "手游内容")[0] == "content"


def test_parse_evidence(tmp_path=None):
    rows = proposer.parse_evidence(EVIDENCE)
    ids = [r["id"] for r in rows]
    assert "performance" in ids and "software" in ids
    perf = next(r for r in rows if r["id"] == "performance")
    assert perf["name"] == "性能表现" and "卡顿" in perf["keywords"]
    assert perf["origin"] == "template:physical"
    soft = next(r for r in rows if r["id"] == "software")
    assert soft["origin"] == "domain"
    assert not any(r["id"].startswith("-") or "---" in r["id"] for r in rows)


def test_build_proposal_merges_template_and_evidence(tmp_path=None):
    ev = proposer.parse_evidence(EVIDENCE)
    proposal = proposer.build_proposal("数码3C", "digital3c", "手机相机电脑",
                                       ["大疆", "华为"], ev)
    ids = [d["id"] for d in proposal["dimensions"]]
    assert len(ids) == len(set(ids))
    assert proposal["template_id"] == "physical"
    assert all(d["origin"] in ("domain",) or d["origin"].startswith("template:")
               for d in proposal["dimensions"])


def test_confirm_sheet_apply_roundtrip(tmp_path: Path):
    proposal = {
        "domain_id": "demo_domain", "name": "演示领域",
        "description": "test", "brands": [],
        "template_id": "physical", "template_reason": ["r"], "template_confidence": 1.0,
        "dimensions": [
            {"id": "quality", "name": "质量", "keywords": ["质量", "做工"],
             "origin": "template:physical", "source": "模板"},
            {"id": "price", "name": "价格", "keywords": ["价格", "贵"],
             "origin": "template:physical", "source": "模板"},
            {"id": "service_x", "name": "售后", "keywords": ["售后", "客服"],
             "origin": "template:physical", "source": "模板"},
        ],
        "risks": [],
    }
    sheet = tmp_path / "confirm.xlsx"
    proposer.write_confirm_sheet(proposal, sheet)
    cache = tmp_path / "cache"
    path = proposer.apply_sheet(sheet, cache, proposal)
    assert path.exists() and path.name == "demo_domain.json"
    schema = json.loads(path.read_text(encoding="utf-8"))
    assert schema["domain_id"] == "demo_domain"
    assert schema["template_id"] == "physical"
    assert len(schema["dimensions"]) == 3
    assert all(d["origin"].startswith("template:") for d in schema["dimensions"])
    # 缺关键词应拒绝
    bad = tmp_path / "bad.xlsx"
    p2 = dict(proposal)
    p2["dimensions"] = [dict(d, keywords=[]) for d in proposal["dimensions"]]
    proposer.write_confirm_sheet(p2, bad)
    try:
        proposer.apply_sheet(bad, cache, p2)
        raised = False
    except ValueError:
        raised = True
    assert raised, "关键词为空应拒绝"


def test_list_domains_includes_cached(tmp_path: Path):
    with mock.patch("app.domains.loader.CACHE_DIR", tmp_path):
        # 写入一个缓存领域
        schema = load_domain("game").model_copy(deep=True)
        schema.domain_id = "cached_x"
        schema.domain_name = "缓存领域"
        (tmp_path / "cached_x.json").write_text(
            json.dumps(schema.model_dump(), ensure_ascii=False), encoding="utf-8")
        from app.domains import loader
        ids = [d["id"] for d in loader.list_domains()]
        assert "cached_x" in ids
        assert all(d.get("cached") for d in loader.list_domains() if d["id"] == "cached_x")


def test_custom_mode_key_baseline_fallback(tmp_path=None):
    """自定义 mode_key 不回落主集基线：baseline_domain_x.json 不存在则返回 None。"""
    from app.core import eval_store
    with mock.patch("app.core.eval_store.FIXTURES_DIR", ROOT / "tests" / "fixtures"):
        assert eval_store.load_frozen_baseline("domain_nonexistent_lexicon") is None
    # 已知主集 key 仍正常（旧格式文件只有 mode，无 mode_key）
    # 2026-08-16：主集词典基线随 golden v1.3（主集维度治理）重测更新为 0.509
    b = eval_store.load_frozen_baseline("lexicon")
    assert b is not None and b.get("accuracy") == 0.509


def test_coldstart_compare_annotations(tmp_path):
    """双 AI 一致性：复合键比对、分歧识别。"""
    from tests import coldstart_annotation as ca
    primary = [
        {"text_id": "a", "text": "t1", "sentiment": "positive", "relevant": "yes",
         "dims": {"功能效果": "positive"}},
        {"text_id": "b", "text": "t2", "sentiment": "neutral", "relevant": "no",
         "dims": {"价格价值": "negative"}},
        {"text_id": "c", "text": "t3", "sentiment": "neutral", "relevant": "no", "dims": {}},
    ]
    secondary = [dict(r) for r in primary]
    rep, disputed = ca.compare_annotations(primary, secondary)
    assert rep["n"] == 3 and rep["sentiment_agreement"] == 1.0 and disputed == []
    secondary[0]["sentiment"] = "negative"
    rep2, dis2 = ca.compare_annotations(primary, secondary)
    assert rep2["sentiment_agreement"] == round(2 / 3, 4) and dis2 == ["a"]


def test_coldstart_finalize_normalization(tmp_path):
    """定版规范归一：relevant=no 清维度/情感 neutral、neutral 强度=1。"""
    import json as _json
    from app.domains import loader as dom_loader
    from tests import coldstart_annotation as ca

    cache = tmp_path / "cache"
    cache.mkdir()
    (cache / "fake.json").write_text(_json.dumps({
        "domain_id": "fake", "domain_name": "假领域", "version": "1.0",
        "template_id": "physical",
        "dimensions": [
            {"id": "quality", "name": "质量", "keywords": ["质量"], "origin": "domain"},
            {"id": "price", "name": "价格", "keywords": ["价格"], "origin": "domain"},
            {"id": "service_x", "name": "售后", "keywords": ["售后"], "origin": "domain"},
        ],
    }, ensure_ascii=False), encoding="utf-8")
    with mock.patch("app.domains.loader.CACHE_DIR", cache):
        sheet = tmp_path / "ann.xlsx"
        ca.gen_worksheet("fake", sheet, samples=[
            {"text_id": "a", "原文": "质量不错", "平台": "demo", "品牌/主题": "X",
             "关键词": "质量", "文本类型": "评论"},
            {"text_id": "b", "原文": "价格太贵", "平台": "demo", "品牌/主题": "X",
             "关键词": "价格", "文本类型": "评论"},
        ])
        from openpyxl import load_workbook
        wb = load_workbook(sheet)
        ws = wb.worksheets[0]
        hdr = [c.value for c in ws[1]]
        i_sent, i_int, i_rel, i_quality, i_price = (
            hdr.index("情感(整条)"), hdr.index("强度(1-5)"), hdr.index("是否相关"),
            hdr.index("质量"), hdr.index("价格"))
        ws.cell(row=2, column=i_sent + 1, value="positive")
        ws.cell(row=2, column=i_int + 1, value="4")
        ws.cell(row=2, column=i_rel + 1, value="yes")
        ws.cell(row=2, column=i_quality + 1, value="positive")
        ws.cell(row=3, column=i_sent + 1, value="negative")
        ws.cell(row=3, column=i_int + 1, value="3")
        ws.cell(row=3, column=i_rel + 1, value="no")
        ws.cell(row=3, column=i_price + 1, value="negative")
        wb.save(sheet)
        out = tmp_path / "golden.csv"
        ca.finalize_merged(sheet, None, out, "fake")
        rows = list(csv.reader(open(out, encoding="utf-8-sig")))
        assert len(rows) == 3  # header + 2
        hdr2 = rows[0]
        b = dict(zip(hdr2, rows[2]))
        assert b["relevant"] == "no" and b["dimension_sentiments"] == "{}"
        assert b["sentiment"] == "neutral" and b["intensity"] == "1"


def main() -> int:
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        fns = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
        failed = 0
        for fn in fns:
            try:
                fn(tmp)
                print(f"✓ {fn.__name__}")
            except AssertionError as exc:
                failed += 1
                print(f"✗ {fn.__name__}: {exc}")
        print(f"\n领域提案引擎测试：{len(fns) - failed}/{len(fns)} 通过")
        return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
