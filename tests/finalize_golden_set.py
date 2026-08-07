# -*- coding: utf-8 -*-
"""黄金集定稿：把已标注的正式标注表 + 复核决议 合并为 golden_set_v1.csv。

用法：
    python tests/finalize_golden_set.py

输入：
    data/datasets/annotation_worksheet_v1.xlsx（游戏标注 / 消费品标注，v0.3 口径）
    data/datasets/cleaning_worksheet_v1.csv（清洗验证集，已标"是否正确丢弃"）

复核决议（2026-08-07，来自 annotation_review_report.xlsx，均为语言现象标记）：
    - v4_C017  移除"反讽"：反讽对象是外部事件而非品牌，品牌情感为真实正向，
               符合 v0.3 "反讽=表面正向/中性、实际负向" 定义
    - v4_C059  追加"黑话"：城墙=官方防御性公关手段（备注已写含义）
    - v2_C096  追加"黑话"：超梦核=网络审美词
    - v4_C014  追加"黑话"：新童鞋=新同学谐音梗
    - v4_P041  不追加"黑话"：不相关文本，语言现象标记对评测无实质影响
"""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

if sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "datasets" / "annotation_worksheet_v1.xlsx"
OUT = ROOT / "data" / "datasets" / "golden_set_v1.csv"
FIXTURE_OUT = ROOT / "tests" / "fixtures" / "golden_set_v1.csv"
CLEANING_SRC = ROOT / "data" / "datasets" / "cleaning_worksheet_v1.csv"
CLEANING_OUT = ROOT / "data" / "datasets" / "golden_set_v1_cleaning.csv"

# text_id -> (操作, 决议说明)
RESOLUTIONS = {
    "v4_C017": ("remove", "复核采纳：移除反讽（反讽对象非品牌，品牌情感为真实正向）"),
    "v4_C059": ("add_slang", "复核采纳：追加黑话（城墙=官方防御性公关手段）"),
    "v2_C096": ("add_slang", "复核采纳：追加黑话（超梦核=网络审美词）"),
    "v4_C014": ("add_slang", "复核采纳：追加黑话（新童鞋=新同学谐音梗）"),
    "v4_P041": ("keep", "复核决议：不追加黑话（不相关文本）"),
}


def parse_flags(raw) -> list[str]:
    if not raw:
        return []
    flags = []
    for x in str(raw).replace("，", ",").replace("、", ",").split(","):
        x = x.strip()
        if x and x not in flags:
            flags.append(x)
    return flags


def main():
    wb = load_workbook(XLSX, read_only=True)
    out_rows = []
    kind_map = {"帖子正文": "post", "评论": "comment", "楼中楼": "reply"}
    for sheet in ("游戏标注", "消费品标注"):
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        hdr = list(next(rows))
        idx = {name: i for i, name in enumerate(hdr)}
        dims = [hdr[j] for j in range(17, 24)]
        for r in rows:
            text_id = r[idx["text_id"]]
            flags = parse_flags(r[idx["语言现象"]])
            note = ""
            if text_id in RESOLUTIONS:
                op, note = RESOLUTIONS[text_id]
                if op == "remove":
                    flags = [f for f in flags if f != "反讽"]
                elif op == "add_slang":
                    if "黑话" not in flags:
                        flags.append("黑话")
            dim_sent = {name: r[17 + i] for i, name in enumerate(dims) if r[17 + i]}
            out_rows.append({
                "text_id": text_id,
                "text": r[idx["原文"]] or "",
                "platform": r[idx["平台"]] or "",
                "domain": r[idx["领域"]] or "",
                "brand": r[idx["品牌/主题"]] or "",
                "keyword": r[idx["关键词"]] or "",
                "kind": kind_map.get(r[idx["文本类型"]] or "", ""),
                "is_reply": "是" if r[idx["是否楼中楼"]] else "",
                "sentiment": r[idx["情感(整条)"]] or "",
                "intensity": r[idx["强度(1-5)"]] or "",
                "dimension_sentiments": json.dumps(dim_sent, ensure_ascii=False),
                "language_flags": ",".join(flags),
                "relevant": r[idx["是否相关"]] or "",
                "remark": r[idx["备注"]] or "",
                "annotator": r[idx["标注人"]] or "",
                "annotated_at": r[idx["标注日期"]] if "标注日期" in idx and r[idx["标注日期"]] else "",
                "url": r[idx["链接"]] or "",
                "time": r[idx["发布时间"]] or "",
                "likes": r[idx["点赞数"]] or "",
                "batch": r[idx["采集批次"]] or "",
                "finalize_note": note,
            })
    wb.close()

    order = {"post": 0, "comment": 1, "reply": 2}
    out_rows.sort(key=lambda x: (x["domain"], x["platform"], order.get(x["kind"], 9), x["text_id"]))
    cols = [
        "text_id", "text", "platform", "domain", "brand", "keyword", "kind", "is_reply",
        "sentiment", "intensity", "dimension_sentiments", "language_flags", "relevant",
        "remark", "annotator", "annotated_at", "url", "time", "likes", "batch", "finalize_note",
    ]
    with open(OUT, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(out_rows)

    # ---- 去 URL 版入库（跨对话/跨机器复现；data/ 在 gitignore） ----
    fixture_cols = [c for c in cols if c not in ("url", "likes", "time")]
    FIXTURE_OUT.parent.mkdir(parents=True, exist_ok=True)
    with open(FIXTURE_OUT, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fixture_cols)
        w.writeheader()
        for r in out_rows:
            w.writerow({c: r[c] for c in fixture_cols})
    print(f"去 URL 版入库：{FIXTURE_OUT.name}，共 {len(out_rows)} 条（不含 url/likes/time）")

    from collections import Counter
    print(f"黄金集定稿：{OUT.name}，共 {len(out_rows)} 条")
    print("情感分布:", dict(Counter(r["sentiment"] for r in out_rows)))
    print("领域分布:", dict(Counter(r["domain"] for r in out_rows)))
    flag_counter: Counter = Counter()
    for r in out_rows:
        if r["language_flags"]:
            for f in r["language_flags"].split(","):
                flag_counter[f] += 1
    print("语言现象标记:", dict(flag_counter))
    print("复核决议应用:", sum(1 for r in out_rows if r["finalize_note"]), "条")

    # ---- 清洗验证集定稿 ----
    from sample_golden_set import normalize_reason

    cleaning_rows = []
    if CLEANING_SRC.exists():
        with open(CLEANING_SRC, encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                brand = Path(str(r["采集批次"] or "")).name
                cleaning_rows.append({
                    "text_id": r["text_id"],
                    "platform": r["平台"] or "",
                    "url": r["链接"] or "",
                    "title": r["标题"] or "",
                    "batch": r["采集批次"] or "",
                    "brand": brand,
                    "reason_raw": r["丢弃原因"] or "",
                    "reason_category": normalize_reason(r["丢弃原因"] or ""),
                    "correct_drop": r["是否正确丢弃"] or "",
                    "annotator": r["标注人"] or "",
                    "annotated_at": r["标注日期"] or "",
                })
    if cleaning_rows:
        cols = [
            "text_id", "platform", "url", "title", "batch", "brand",
            "reason_raw", "reason_category", "correct_drop", "annotator", "annotated_at",
        ]
        with open(CLEANING_OUT, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=cols)
            w.writeheader()
            w.writerows(cleaning_rows)
        confirmed = sum(1 for r in cleaning_rows if r["correct_drop"] == "应丢弃")
        print(f"清洗验证集定稿：{CLEANING_OUT.name}，共 {len(cleaning_rows)} 条，确认应丢弃 {confirmed} 条")
    else:
        print("清洗验证集未标注，跳过定稿")


if __name__ == "__main__":
    main()
