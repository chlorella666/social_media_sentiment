"""内置演示报告单测：离线生成、结构完整、数字与证据自洽。"""

from __future__ import annotations

import sys
import unittest
import io
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app import demo_report


class TestDemoReport(unittest.TestCase):
    def setUp(self):
        # 测试强制使用内置虚构「云朵咖啡」源（本地可能存在真实演示源，避免断言漂移）
        self._orig_source = demo_report._DEMO_SOURCE
        demo_report._DEMO_SOURCE = (
            ROOT / "data" / "state" / "demo_report" / "__test_nonexistent__.json"
        )

    def tearDown(self):
        demo_report._DEMO_SOURCE = self._orig_source

    def test_build_demo_bundle_structure(self):
        b = demo_report.build_demo_bundle()
        s = b.summary
        self.assertGreaterEqual(s["total_items"], 100)  # 演示口径：样本充足
        self.assertEqual(s["total_posts"], 26)
        self.assertEqual(len(b.findings), 5)
        self.assertGreaterEqual(len(b.evidence), 10)
        self.assertEqual(b.insight_mode, "llm")
        self.assertTrue(
            all(it.method == "llm" for it in b.coded_items),
            "演示报告应全部为 LLM 精分析编码（展示 LLM 模式成品效果）",
        )
        self.assertEqual(b.warnings, [], "演示报告不应出现质量警示")
        # F-018（2026-08-26，修订版回退）：LLM 模式不再产出 structured_summary
        self.assertEqual(b.structured_summary, {}, "LLM 模式不应有 structured_summary")
        self.assertEqual(b.structured_summary_source, "rule")
        self.assertGreaterEqual(
            s.get("narrative_stats", {}).get("total", 0), 10,
            "叙事/归因样本应足以展示聚合图",
        )
        self.assertGreaterEqual(len(s.get("dimensions") or {}), 4)

    def test_llm_no_structured_summary_zone(self):
        """F-018（修订版回退）：LLM 模式演示报告四端不渲染 structured_summary 区。"""
        import io as _io

        from app.output.excel_writer import build_excel
        from app.output.html_report import build_html
        from app.output.word_report import build_word

        b = demo_report.build_demo_bundle()
        html = build_html(b)
        self.assertNotIn("解读与建议", html, "LLM 模式 HTML 不应有 structured_summary 区")
        self.assertIn("核心发现", html, "LLM 模式应保留 findings（深度发现+行动）")
        wb_bytes = build_excel(b).getvalue()
        import openpyxl
        wb = openpyxl.load_workbook(_io.BytesIO(wb_bytes))
        self.assertNotIn("解读与建议", wb.sheetnames, "LLM 模式 Excel 不应有解读 sheet")
        doc_bytes = build_word(b).getvalue()
        from docx import Document
        doc = Document(_io.BytesIO(doc_bytes))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("六、解读与建议", text, "LLM 模式 Word 不应有解读段")

    def test_builtin_demo_source(self):
        """F-020：内置 app/demo_source.json 优先加载（恋与深空脱敏样本 + LLM topics）。"""
        import json as _json

        builtin = Path(__file__).resolve().parent.parent / "app" / "demo_source.json"
        self.assertTrue(builtin.exists(), "缺少 app/demo_source.json")
        self._orig_source = demo_report._DEMO_SOURCE
        demo_report._DEMO_SOURCE = None
        try:
            b = demo_report.build_demo_bundle()
            s = b.summary
            self.assertGreaterEqual(s["total_items"], 200, "演示数据应足量（恋与深空 304 条）")
            self.assertEqual(b.insight_mode, "llm")
            self.assertGreaterEqual(len(s.get("topics") or []), 5, "LLM 编码工作流 topics 应固化")
            self.assertFalse(b.structured_summary, "LLM 模式无 structured_summary")
            raw = _json.loads(builtin.read_text(encoding="utf-8"))
            self.assertEqual(raw["_demo_meta"]["source"], "恋与深空玩家讨论（脱敏样本）")
        finally:
            demo_report._DEMO_SOURCE = self._orig_source

    def test_v015_layout_assertions(self):
        """F-036：报告布局（分析结论前置 / 情感矩阵树 / 方法说明折叠 / 查询串折叠 / 无「需复核」）。"""
        import io as _io

        from app.output.html_report import build_html

        b = demo_report.build_demo_bundle()
        html = build_html(b)
        # 分析结论在核心发现之前
        i_conclusion = html.find("分析结论")
        i_findings = html.find("核心发现")
        self.assertGreater(i_conclusion, 0)
        self.assertGreater(i_findings, 0)
        self.assertLess(i_conclusion, i_findings, "分析结论应前置")
        # 情感矩阵树（F-035 treemap）
        self.assertIn("情感矩阵树", html)
        self.assertIn("h3", html)
        # F-037：主题观点折叠（details）+ 四处折叠齐全
        self.assertIn("🧩 主题观点（代表观点）", html)
        self.assertGreaterEqual(html.count("<details"), 3, "主题观点/方法说明/查询串应折叠")
        # 方法与数据说明折叠表格
        self.assertIn("数据计算方法（点击展开）", html)
        self.assertIn("<details>", html)
        # 实际查询串折叠
        self.assertIn("实际查询串（按渠道）", html)
        # 演示报告无「需复核」样本提示（方法说明折叠表格中的口径说明属正常）
        self.assertNotIn("条样本待人工复核", html)
        # need_review 字段清零
        self.assertEqual(sum(1 for it in b.coded_items if it.need_review), 0)
        print("F-036 report layout assertions passed")

    def test_evidence_refs_all_resolve(self):
        b = demo_report.build_demo_bundle()
        ids = {e["id"] for e in b.evidence}
        missing = [
            r for f in b.findings for r in (f.get("evidence_refs") or [])
            if r not in ids
        ]
        self.assertEqual(missing, [], "发现的证据引用必须能在证据卡中找到")

    def test_findings_match_summary_numbers(self):
        b = demo_report.build_demo_bundle()
        s = b.summary
        dist = s["sentiment_distribution"]
        f1 = next(f for f in b.findings if f["id"] == "F1")
        self.assertIn(f"共 {s['total_items']} 条", f1["claim"])
        self.assertIn(f"正面 {dist['positive']['count']} 条", f1["claim"])
        # F2/F3 的维度负面率与 summary 一致
        dims = s.get("dimensions") or {}
        for fid, dim in (("F2", "channel_service"), ("F3", "price_value")):
            f = next(x for x in b.findings if x["id"] == fid)
            rate = dims[dim]["negative_rate"] * 100
            self.assertIn(f"{rate:.0f}%", f["claim"])

    def test_build_is_deterministic(self):
        a = demo_report.build_demo_bundle()
        c = demo_report.build_demo_bundle()
        self.assertEqual(
            a.summary["sentiment_distribution"],
            c.summary["sentiment_distribution"],
        )
        self.assertEqual(
            [f["claim"] for f in a.findings],
            [f["claim"] for f in c.findings],
        )

    def test_accuracy_compare_md(self):
        md = demo_report.accuracy_compare_md()
        self.assertIn("词典模式", md)
        self.assertIn("LLM 精分析", md)
        self.assertIn("47%~51%", md)
        self.assertIn("80%~88%", md)

    def test_prepare_demo_files(self):
        b = demo_report.build_demo_bundle()
        files = demo_report.prepare_demo_files(b)
        self.assertTrue(files["excel"], "Excel 产物不应为空")
        self.assertIn("report", files["html"], "HTML 产物不应为空")
        self.assertTrue(demo_report.DEMO_RESULT.exists())

    def test_excel_export_coded_rows_and_no_formula(self):
        """Excel 导出：情感编码明细行数与编码条数一致，且无 '=' 开头单元格
        （2026-08-21 修复：_coded_rows 缩进 bug 致明细空白；'== 表头 ==' 被
        Excel 当作公式触发修复提示）。"""
        import openpyxl

        from app.output.excel_writer import build_excel

        b = demo_report.build_demo_bundle()
        buf = build_excel(b)
        wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
        ws = wb["情感编码明细"]
        self.assertEqual(
            ws.max_row - 1, len(b.coded_items),
            "情感编码明细行数应与编码条数一致",
        )
        self.assertGreaterEqual(ws.max_column, 17)
        formula_cells = [
            c.coordinate
            for sheet in wb.worksheets
            for row in sheet.iter_rows()
            for c in row
            if isinstance(c.value, str) and c.value.startswith("=")
        ]
        self.assertEqual(formula_cells, [], "不应有 '=' 开头单元格（Excel 公式注入）")

    def test_ui_quickstart_opens_demo_report(self):
        """端到端：快速上手卡「✨ 看演示报告」→ 结果页（stage 5），无异常。"""
        from streamlit.testing.v1 import AppTest

        at = AppTest.from_file(
            str(ROOT / "app" / "main.py"), default_timeout=90
        )
        at.run()
        # 1.5 首次确认门禁：全新环境（CI）只渲染《使用边界》确认页，
        # 先勾选并确认再进入应用（已确认环境无复选框，自动跳过）
        agree = [c for c in at.checkbox if c.label.startswith("我已阅读并同意")]
        if agree:
            agree[0].set_value(True).run()
            confirm = next(
                (b for b in at.button if b.label == "确认并进入应用"), None
            )
            self.assertIsNotNone(confirm, "首次确认页应有「确认并进入应用」按钮")
            confirm.click().run()
        button_labels = [b.label for b in at.button]
        self.assertNotIn("▶️ 一键体验", button_labels, "「一键体验」按钮应已移除")
        marks = " ".join(str(m.value) for m in at.markdown)
        self.assertNotIn(
            "词典模式 vs LLM 模式：准确率对比", marks,
            "侧边栏准确率对比折叠区应已移除（保留开关下方常驻提示）",
        )
        target = next(
            (b for b in at.button if "看演示报告" in b.label), None
        )
        self.assertIsNotNone(target, "快速上手卡应包含「看演示报告」按钮")
        target.click().run()
        self.assertEqual(at.session_state["stage"], 5)
        self.assertTrue(at.session_state["demo_report"])
        self.assertIn("bundle", at.session_state)
        self.assertIsNotNone(at.session_state["bundle"])
        self.assertEqual(at.exception, [], "演示报告结果页不应有异常")
        self.assertTrue(
            any("⑥ 分析结果" == s.value for s in at.subheader),
            "结果页应渲染分析结果区",
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)
