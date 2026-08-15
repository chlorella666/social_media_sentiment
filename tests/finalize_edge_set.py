# -*- coding: utf-8 -*-
"""边界样本专项集标注编排（2.3 阶段 1 配套）。

用法：
    python tests/finalize_edge_set.py --worksheets   # 生成主标/副标空白标注表（xlsx+CSV）
    python tests/finalize_edge_set.py --compare      # 合并双 AI 结果：一致性报告 + 两张复核表
    python tests/finalize_edge_set.py                # 打印用法

流程（docs/边界样本专项集方案.md §四）：
    工作表生成 → 双 AI 标注（workbuddy 主标 / trae-GLM-5.2 副标）→
    本脚本 --compare 输出一致性报告与复核表 → 人工填复核表 →
    tests/apply_edge_review.py 定版。
"""

from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from tests import edge_annotation as ea  # noqa: E402


def _pick_input(primary: bool) -> Path:
    x = ea.EDGE_PRIMARY_XLSX if primary else ea.EDGE_SECONDARY_XLSX
    c = ea.EDGE_PRIMARY_CSV if primary else ea.EDGE_SECONDARY_CSV
    return ea.pick_annotated_input(x, c)


def _filled_count(path: Path) -> int:
    try:
        return sum(1 for r in ea.load_annotation(path) if r["sentiment"])
    except Exception:
        return -1


def gen_worksheets(force: bool = False, allow_annotated: bool = False) -> None:
    """生成主标/副标两份空白标注表（xlsx 含下拉校验，CSV 便于程序化标注）。

    安全闸门（2026-08-13 事故后新增）：--force 覆盖已标注文件会静默丢失标注
    （品牌元数据修复曾因此覆盖 workbuddy/trae 已填写的 300 条），故已填情感>0 的
    目标文件必须显式传入 allow_annotated=True 才允许覆盖。
    """
    if not ea.EDGE_ANNOT_XLSX.exists():
        raise SystemExit(f"基础标注表不存在：{ea.EDGE_ANNOT_XLSX}（先运行 tests/sample_golden_set.py --edge-only）")
    for out in (ea.EDGE_PRIMARY_XLSX, ea.EDGE_SECONDARY_XLSX,
                ea.EDGE_PRIMARY_CSV, ea.EDGE_SECONDARY_CSV):
        if out.exists():
            filled = _filled_count(out)
            if not force:
                print(f"跳过（已存在，用 --force 覆盖）：{out.name}")
                continue
            if filled > 0 and not allow_annotated:
                raise SystemExit(
                    f"拒绝覆盖已标注文件：{out.name}（情感已填 {filled} 条）。"
                    "如需确认覆盖请加 --allow-annotated（会导致已填标注丢失）"
                )
        if out.suffix == ".xlsx":
            shutil.copy2(ea.EDGE_ANNOT_XLSX, out)
        else:
            shutil.copy2(ea.EDGE_ANNOT_CSV, out)
        print(f"已生成：{out}")


def write_review_xlsx(path: Path, rows: list[dict], title: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "复核"
    header_fill = PatternFill("solid", fgColor="E2EFDA")
    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    ws.append(ea.REVIEW_COLUMNS)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    for r in rows:
        ws.append([r.get(c, "") for c in ea.REVIEW_COLUMNS])
    widths = [6, 14, 50, 10, 8, 14, 10, 10, 10, 8, 8, 16, 16, 8, 8, 12, 12, 18,
              10, 8, 16, 8, 12, 20, 10, 12]
    for i, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    ws.freeze_panes = "C2"

    guide = wb.create_sheet("填写说明")
    lines = [
        title,
        "",
        "一、对照「主标」与「副标」两栏（情感/强度/维度/相关/语言现象），在「人工判定*」列填写最终值：",
        "    - 同意某一方 → 填该方的取值；两方都不同意 → 填你认为正确的值；",
        "    - 判定列留空 = 认可主标，无需填写；",
        "    - 维度列填写 JSON（如 {\"价格价值\":\"negative\"}）或「维度:值」形式；",
        "二、「是否相关」用 是/否（跑题/引流填 否，整条 neutral，不计入主评分）；",
        "三、「备注」可写判定理由（分歧样本建议写，作为争议集弹药）；",
        "四、填完后运行 tests/apply_edge_review.py 定版。",
    ]
    for line in lines:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    guide["A1"].font = Font(bold=True, size=12)
    wb.save(path)
    print(f"已生成：{path}（{len(rows)} 条）")


def run_compare() -> int:
    primary_path = _pick_input(True)
    secondary_path = _pick_input(False)
    if not primary_path.exists() or not secondary_path.exists():
        print("主标/副标标注表未就绪：")
        print(f"  主标：{ea.EDGE_PRIMARY_XLSX} 或 {ea.EDGE_PRIMARY_CSV}（{primary_path.exists()}）")
        print(f"  副标：{ea.EDGE_SECONDARY_XLSX} 或 {ea.EDGE_SECONDARY_CSV}（{secondary_path.exists()}）")
        print("先用 --worksheets 生成空白表，标注完成后重试。")
        return 1

    primary = ea.load_annotation(primary_path)
    secondary = ea.load_annotation(secondary_path)
    p_ids = {r["text_id"] for r in primary}
    s_ids = {r["text_id"] for r in secondary}
    print(f"主标 {len(primary)} 条 / 副标 {len(secondary)} 条；"
          f"主标缺失 {len(s_ids - p_ids)} / 副标缺失 {len(p_ids - s_ids)}")

    p_filled = sum(1 for r in primary if r["sentiment"])
    s_filled = sum(1 for r in secondary if r["sentiment"])
    if p_filled != len(primary) or s_filled != len(secondary):
        print(f"标注未完成：主标已填情感 {p_filled}/{len(primary)}，副标 {s_filled}/{len(secondary)}"
              "（防止在空白/半成品表上误生成虚假一致率报告）")
        return 1

    errs = ea.validate_rows(primary) + ea.validate_rows(secondary)
    if errs:
        print("取值校验失败（不生成复核表）：")
        for e in errs[:30]:
            print("  -", e)
        print(f"共 {len(errs)} 条错误")
        return 1

    report = ea.merge_compare(primary, secondary)
    consistency_ids, calibration_ids = ea.build_review_sheets(primary, secondary, report)
    p = {r["text_id"]: r for r in primary}
    s = {r["text_id"]: r for r in secondary}
    consistency_rows = [ea.review_row(i, tid, p, s) for i, tid in enumerate(consistency_ids, 1)]
    calibration_rows = [ea.review_row(i, tid, p, s) for i, tid in enumerate(calibration_ids, 1)]

    ea.EDGE_CONSISTENCY_REPORT.parent.mkdir(parents=True, exist_ok=True)
    ea.EDGE_CONSISTENCY_REPORT.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    write_review_xlsx(ea.EDGE_REVIEW_CONSISTENCY_XLSX, consistency_rows,
                      "边界集一致性核对表（随机 20% + 高歧义 100%）")
    write_review_xlsx(ea.EDGE_REVIEW_CALIBRATION_XLSX, calibration_rows,
                      "边界集人工抽检校准表（60 条，seed 可复现）")

    print(f"\n一致性报告：{ea.EDGE_CONSISTENCY_REPORT}")
    print(f"整条情感一致率：{report['sentiment_agreement']:.1%}（目标 ≥80%）"
          f"｜kappa {report['kappa']:.3f}（目标 ≥0.6）")
    if report["dimension_agreement"] is not None:
        print(f"维度级一致率：{report['dimension_agreement']:.1%}（n={report['dimension_n']}，目标 ≥75%）"
              f"｜相关一致率 {report['relevance_agreement']:.1%}")
    else:
        print(f"维度级一致率：无维度标注（n={report['dimension_n']}）"
              f"｜相关一致率 {report['relevance_agreement']:.1%}")
    print("按子集一致率：", {
        k: f"{v['agreement']:.0%}(n={v['n']})" for k, v in sorted(report["by_subset"].items())
    })
    print(f"高歧义 {len(report['high_ambiguity'])} 条（反讽/mixed/relevant=no）"
          f"｜分歧 {len(report['disputes'])} 条")
    return 0


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    import argparse

    ap = argparse.ArgumentParser(description="边界样本专项集标注编排")
    ap.add_argument("--worksheets", action="store_true", help="生成主标/副标空白标注表")
    ap.add_argument("--compare", action="store_true", help="合并双 AI 结果并生成复核表")
    ap.add_argument("--acceptance-review", action="store_true",
                    help="生成模型错误人工验收表（V2 验证纪律；--blind 为盲审档）")
    ap.add_argument("--report", type=Path, help="benchmark 报告 JSON（--acceptance-review 用）")
    ap.add_argument("--out", type=Path, help="输出 xlsx（--acceptance-review 用）")
    ap.add_argument("--blind", action="store_true",
                    help="盲审：验收表省略模型判定列")
    ap.add_argument("--force", action="store_true", help="覆盖已存在的产物")
    ap.add_argument("--allow-annotated", dest="allow_annotated", action="store_true",
                    help="（危险）允许覆盖已填写的标注表")
    args = ap.parse_args()

    if not (args.worksheets or args.compare or args.acceptance_review):
        ap.print_help()
        return 1
    if args.worksheets:
        gen_worksheets(force=args.force, allow_annotated=args.allow_annotated)
    if args.compare:
        return run_compare()
    if args.acceptance_review:
        if not (args.report and args.out):
            ap.error("--acceptance-review 需要 --report 与 --out")
        report = json.loads(args.report.read_text(encoding="utf-8"))
        errors = report.get("errors") or []
        ea.gen_acceptance_review_xlsx(errors, args.out, blind=args.blind)
        return 0
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
