# -*- coding: utf-8 -*-
"""边界集「品牌/主题」元数据修复（2.3 阶段 1 补，方案 A 已确认）。

背景：反讽/方言/emoji 子集由"现象信号词"采集（唔该/巴适/😡😡 太坑了 等），
采样器把采集关键词原样写入「品牌/主题」列，导致 120 行出现伪品牌；且整批现象
采集被 collect_edge_samples 统一标为 consumer 领域（仅影响维度列选择）。
影响：标注者无法按"是否与品牌/主题相关"口径判定。

方案 A（已确认）：现象样本「品牌/主题」置空（关键词列保留真实搜索词），
说明页补现象样本判定指引；领域列保持原值（仅用于选择维度列，不适用则留空）。

本脚本为**原地列级修复**：不重采样、不改 text_id/子集/渠道/原文/关键词，
只清空命中现象关键词的「品牌/主题」列，并校验前后一致性。

用法：
    python tests/fix_edge_brand_meta.py --check   # 只检查命中行数与分布，不改文件
    python tests/fix_edge_brand_meta.py           # 执行修复 + 重建主标/副标空白表
"""

from __future__ import annotations

import argparse
import csv
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from tests import collect_edge_samples as ces  # noqa: E402
from tests import edge_annotation as ea  # noqa: E402

# 现象采集脚本的关键词表（与内置方言/反讽/emoji 信号词合并判定）
EXTRA_KEYWORDS = frozenset(
    kw for cfg in ces.EDGE_PLAN.values() for kw in cfg["keywords"]
)

# 说明页追加的现象样本判定指引（幂等：已包含则跳过）
GUIDE_LINES = [
    "",
    "六、现象样本（品牌/主题 为空）：反讽/方言/emoji 等子集由现象信号词采集",
    "    （如「唔该」「😡😡 太坑了」），不是真实品牌。此时「是否相关」判定口径为：",
    "    是否为可分析的真实内容（非页面壳/导航/纯系统提示/纯转发）；有情感即可判定",
    "    （neutral 也算），与品牌无关；维度不适用时全部留空，整条按文本情感标。",
    "七、领域列仅用于选择维度列（游戏/消费品）；现象样本若内容与领域无关，维度留空即可。",
]


def _load_csv(path: Path) -> tuple[list[dict], list[str]]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    return rows, list(rows[0].keys()) if rows else []


def matched_rows(path: Path) -> list[dict]:
    rows, _ = _load_csv(path)
    return [r for r in rows if ea.is_phenomenon_brand(r["品牌/主题"], EXTRA_KEYWORDS)]


def _snapshot(rows: list[dict]) -> dict[str, tuple]:
    return {
        r["text_id"]: (r["关键词"], r["子集"], r["平台"], r["原文"], r["领域"])
        for r in rows
    }


def fix_csv(path: Path) -> list[str]:
    rows, fields = _load_csv(path)
    before = _snapshot(rows)
    changed = []
    for r in rows:
        if ea.is_phenomenon_brand(r["品牌/主题"], EXTRA_KEYWORDS):
            changed.append(r["text_id"])
            r["品牌/主题"] = ""
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)
    verify_csv(path, before, set(changed))
    return changed


def verify_csv(path: Path, before: dict, changed_ids: set[str]) -> None:
    rows, _ = _load_csv(path)
    after = _snapshot(rows)
    assert set(after) == set(before), "text_id 集合变化，禁止修复"
    for tid in before:
        if tid not in after:
            raise AssertionError(f"text_id 丢失：{tid}")
        if after[tid] != before[tid]:
            raise AssertionError(f"非品牌列被改动：{tid}")
    for r in rows:
        brand = r["品牌/主题"]
        if r["text_id"] in changed_ids:
            assert brand == "", f"{r['text_id']} 品牌未置空：{brand!r}"
        else:
            assert brand, f"{r['text_id']} 非现象样本品牌被清空"


def fix_xlsx(path: Path, changed_ids: set[str]) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    for sheet in ("游戏标注", "消费品标注"):
        ws = wb[sheet]
        hdr = [c.value for c in ws[1]]
        i_brand = hdr.index("品牌/主题") + 1
        for row in ws.iter_rows(min_row=2):
            tid = row[1].value
            if tid and str(tid).strip() in changed_ids:
                row[i_brand - 1].value = ""
    guide = wb["说明"]
    lines = [str(r[0]) for r in guide.iter_rows(values_only=True) if r[0]]
    if not any("现象样本" in line for line in lines):
        for line in GUIDE_LINES:
            guide.append([line])
    wb.save(path)


def verify_xlsx(path: Path, changed_ids: set[str]) -> int:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    checked = 0
    try:
        for sheet in ("游戏标注", "消费品标注"):
            ws = wb[sheet]
            rows = ws.iter_rows(values_only=True)
            hdr = list(next(rows))
            i_brand = hdr.index("品牌/主题")
            for r in rows:
                if not r or not r[1]:
                    continue
                tid = str(r[1]).strip()
                if tid in changed_ids:
                    assert r[i_brand] in ("", None), f"xlsx {tid} 品牌未置空"
                    checked += 1
                else:
                    assert r[i_brand], f"xlsx {tid} 非现象样本品牌被清空"
    finally:
        wb.close()
    return checked


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="边界集「品牌/主题」元数据修复（方案 A）")
    ap.add_argument("--check", action="store_true", help="只检查命中行数，不改文件")
    args = ap.parse_args()

    base_csv, base_xlsx = ea.EDGE_ANNOT_CSV, ea.EDGE_ANNOT_XLSX
    matched = matched_rows(base_csv)
    print(f"命中现象样本：{len(matched)} 行（共 {len(_load_csv(base_csv)[0])} 行）")
    print("按子集：", dict(Counter(r["子集"] for r in matched)))
    print("按领域：", dict(Counter(r["领域"] for r in matched)))
    brands = Counter(r["品牌/主题"] for r in matched)
    print("现象关键词样例：", {k: v for k, v in brands.most_common(12)})

    if args.check:
        print("（--check 模式，未改动任何文件）")
        return 0

    if not matched:
        print("没有需要修复的行，直接重建主标/副标空白表。")
    changed_ids = set(fix_csv(base_csv))
    print(f"CSV 已修复：{len(changed_ids)} 行品牌置空，其余列与 text_id 集合校验通过")

    fix_xlsx(base_xlsx, changed_ids)
    checked = verify_xlsx(base_xlsx, changed_ids)
    print(f"XLSX 已修复并校验：{checked} 行品牌置空，说明页已补现象样本指引")

    from tests import finalize_edge_set as fes

    fes.gen_worksheets(force=True)
    print("主标/副标空白表已从修复后的基表重建")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
