# -*- coding: utf-8 -*-
"""2.6 验证纪律工程测试（V1 hold-out / V2 gold 修订独立）。

覆盖：
- V1 hold-out 分层切分（渠道/品牌配额 + 时间窗三分桶）、数量钳制；
- V1 冻结只读守卫 + 使用次数上限（≤2，重切 --force 清零）；
- V2 benchmark --exclude-gold-revised-by-model 过滤；
- V2 验收表生成/读取（--blind 盲审档、是否参考过模型判定列）。
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from tests import benchmark_golden as bg  # noqa: E402
from tests import coldstart_annotation as cs  # noqa: E402
from tests import edge_annotation as ea  # noqa: E402


def _pool(n_per_key: int = 40) -> list[dict]:
    platforms = ["weibo", "websearch", "xiaohongshu"]
    brands = ["大疆", "影石"]
    dates = ["2026-08-01", "2026-08-05", "2026-08-09"]
    items = []
    i = 0
    for pl in platforms:
        for br in brands:
            for d in dates:
                for k in range(n_per_key):
                    i += 1
                    items.append({
                        "text_id": f"t{i:04d}",
                        "原文": f"{pl}-{br}-{d}-{k}",
                        "平台": pl,
                        "品牌/主题": br,
                        "关键词": br,
                        "文本类型": "评论",
                        "pub_date": d,
                    })
    return items


def test_holdout_split_quotas_and_time_window() -> None:
    pool = _pool(40)
    holdout, iteration, rep = cs.split_holdout(pool, 30, seed=42)
    assert len(holdout) == 30
    assert len(iteration) == len(pool) - 30
    assert rep["issues"] == []
    # 时间窗三分桶全部覆盖
    dates = [d for d in (cs._pub_date(i["pub_date"]) for i in pool) if d]
    buckets = {cs._time_bucket(r.get("pub_date"), min(dates), max(dates))
               for r in holdout}
    assert buckets == {"0", "1", "2"}
    # 同分布：各 (平台, 品牌) 在 hold-out 中的占比与池一致（无 issue 即校验通过）
    by_key = {}
    for r in pool:
        by_key[(r["平台"], r["品牌/主题"])] = by_key.get((r["平台"], r["品牌/主题"]), 0) + 1
    hold_keys = {}
    for r in holdout:
        hold_keys[(r["平台"], r["品牌/主题"])] = hold_keys.get((r["平台"], r["品牌/主题"]), 0) + 1
    for k, n in by_key.items():
        assert abs(hold_keys.get(k, 0) / len(holdout) - n / len(pool)) <= 0.11
    print("✓ hold-out 分层切分（配额 + 时间窗覆盖）通过")


def test_holdout_split_shortfall_clamp() -> None:
    pool = []
    for pl in ("weibo", "websearch"):
        for i in range(4):
            pool.append({
                "text_id": f"t{len(pool) + 1:04d}",
                "原文": f"{pl}-s{i}",
                "平台": pl,
                "品牌/主题": "大疆",
                "关键词": "大疆",
                "文本类型": "评论",
                "pub_date": "2026-08-01",
            })
    holdout, iteration, _ = cs.split_holdout(pool, 30, seed=42)
    assert len(pool) == 8 and len(holdout) == 8
    assert iteration == []
    print("✓ hold-out 数量钳制（池小于目标时全取）通过")


def test_holdout_split_missing_dates() -> None:
    """缺 pub_date 的样本归 'unknown' 桶，不抛错且照常分层。"""
    pool = []
    for i in range(6):
        pool.append({
            "text_id": f"d{i:04d}", "原文": f"dated-{i}",
            "平台": "weibo", "品牌/主题": "大疆",
            "关键词": "大疆", "文本类型": "评论", "pub_date": "2026-08-01",
        })
    for i in range(6):
        pool.append({
            "text_id": f"u{i:04d}", "原文": f"undated-{i}",
            "平台": "weibo", "品牌/主题": "大疆",
            "关键词": "大疆", "文本类型": "评论", "pub_date": "",
        })
    holdout, iteration, rep = cs.split_holdout(pool, 6, seed=42)
    assert len(holdout) == 6 and len(iteration) == 6
    assert rep["issues"] == []
    assert "unknown" in {cs._time_bucket(r.get("pub_date"), None, None)
                         for r in holdout}
    print("✓ hold-out 缺失日期样本（unknown 桶）通过")


def test_holdout_freeze_and_use_guard() -> None:
    tmp = Path(tempfile.mkdtemp(prefix="sms_holdout_"))
    golden = tmp / "holdout_demo.csv"
    golden.write_text("text_id,sentiment\nh1,positive\nh2,negative\n", encoding="utf-8-sig")
    marker = cs.freeze_holdout(golden)
    assert marker.exists()
    # 冻结后只读：重复冻结拒绝（除非 --force）
    try:
        cs.freeze_holdout(golden)
        raise AssertionError("应拒绝重复冻结")
    except ValueError:
        pass
    # 使用次数 ≤2
    r1 = cs.record_holdout_use(golden)
    r2 = cs.record_holdout_use(golden)
    assert r1["ok"] and r2["ok"] and r2["uses"] == 2
    r3 = cs.record_holdout_use(golden)
    assert r3["ok"] is False and r3.get("blocked")
    # 重切（--force）作废旧卷并清零
    cs.freeze_holdout(golden, force=True)
    assert cs.record_holdout_use(golden)["uses"] == 1
    # 冻结后改动内容 → 运行时指纹校验拦截（卷子作废）
    golden.write_text("text_id,sentiment\nh1,positive\nh2,negative\nh3,neutral\n",
                      encoding="utf-8-sig")
    res = cs.record_holdout_use(golden)
    assert res["ok"] is False and res.get("blocked")
    assert "指纹" in res.get("error", "")
    # 未冻结的 golden → 拒绝记录
    unfrozen = tmp / "holdout_not_frozen.csv"
    unfrozen.write_text("text_id,sentiment\nh1,positive\n", encoding="utf-8-sig")
    assert cs.record_holdout_use(unfrozen)["ok"] is False
    print("✓ hold-out 冻结只读 + 使用上限（≤2 + force 重切）通过")


def test_gold_revised_filter() -> None:
    rows = [
        {"text_id": "a", "gold_revised_by_model": "是"},
        {"text_id": "b", "gold_revised_by_model": "否"},
        {"text_id": "c", "gold_revised_by_model": ""},
        {"text_id": "d", "gold_revised_by_model": "yes"},
    ]
    assert bg.is_gold_revised_by_model("是") and bg.is_gold_revised_by_model("Yes")
    assert not bg.is_gold_revised_by_model("否") and not bg.is_gold_revised_by_model("")
    kept, dropped = bg.filter_gold_revised(rows, True)
    assert {r["text_id"] for r in dropped} == {"a", "d"}
    assert {r["text_id"] for r in kept} == {"b", "c"}
    kept2, dropped2 = bg.filter_gold_revised(rows, False)
    assert len(kept2) == 4 and dropped2 == []
    print("✓ gold 修订独立过滤（--exclude-gold-revised-by-model）通过")


def test_acceptance_review_sheet_roundtrip() -> None:
    assert "是否参考过模型判定" in ea.REVIEW_COLUMNS
    errors = [{
        "text_id": "t1", "text": "原文", "platform": "weibo", "domain": "game",
        "subset": "irony", "gold": "negative", "pred": "positive", "llm_used": True,
    }]
    tmp = Path(tempfile.mkdtemp(prefix="sms_accept_"))
    out = tmp / "acceptance.xlsx"
    ea.gen_acceptance_review_xlsx(errors, out, blind=False)
    rows = ea.load_acceptance_xlsx(out)
    assert rows[0]["text_id"] == "t1"
    assert rows[0]["pred"] == "positive"
    assert rows[0]["human_sentiment"] == ""
    assert "gold_revised_by_model" in rows[0]
    out_blind = tmp / "acceptance_blind.xlsx"
    ea.gen_acceptance_review_xlsx(errors, out_blind, blind=True)
    rows_b = ea.load_acceptance_xlsx(out_blind)
    assert rows_b[0]["pred"] == ""  # 盲审档省略模型判定列
    assert rows_b[0]["text_id"] == "t1"
    print("✓ 验收表生成/读取（含 --blind 盲审）通过")


def test_load_human_decisions() -> None:
    """人工复核表回填解析：text_id → {sentiment, relevant}；留空不覆盖。"""
    from openpyxl import Workbook

    tmp = Path(tempfile.mkdtemp(prefix="sms_human_"))
    p = tmp / "review.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["text_id", "原文", "人工判定情感", "人工判定相关"])
    ws.append(["h1", "文本一", "neutral", "no"])
    ws.append(["h2", "文本二", "", "yes"])
    ws.append(["h3", "文本三", "positive", ""])
    wb.save(p)
    out = cs.load_human_decisions(p)
    assert out["h1"] == {"sentiment": "neutral", "relevant": "no"}
    assert out["h2"] == {"sentiment": "", "relevant": "yes"}
    assert out["h3"] == {"sentiment": "positive", "relevant": ""}
    assert "h4" not in out
    print("✓ 人工复核表回填解析 通过")


def test_load_human_decisions_shared_text_id() -> None:
    """同帖多评论共用 text_id：用 (text_id, 原文) 复合键，避免误套。"""
    from openpyxl import Workbook

    tmp = Path(tempfile.mkdtemp(prefix="sms_human_dup_"))
    p = tmp / "review_dup.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["text_id", "原文", "人工判定情感", "人工判定相关"])
    ws.append(["P1:comment", "评论一", "neutral", "yes"])
    ws.append(["P1:comment", "评论二", "positive", "no"])
    ws.append(["P2:comment", "唯一评论", "negative", "yes"])
    wb.save(p)
    out = cs.load_human_decisions(p)
    assert ("P1:comment", "评论一") in out
    assert ("P1:comment", "评论二") in out
    assert "P1:comment" not in out  # 共用 text_id 不再用裸 id 键
    assert out["P2:comment"]["sentiment"] == "negative"
    print("✓ 同帖多评论复合键回填 通过")


def test_gen_dispute_review_roundtrip() -> None:
    """分歧复核表：只收情感/相关分歧，维度分歧不进入；可回填解析。"""
    tmp = Path(tempfile.mkdtemp(prefix="sms_dispute_"))
    primary = [
        {"text_id": "a", "text": "t1", "platform": "weibo", "brand": "大疆",
         "sentiment": "neutral", "relevant": "yes", "dims": {"功能效果": "positive"}},
        {"text_id": "b", "text": "t2", "platform": "bilibili", "brand": "影石",
         "sentiment": "negative", "relevant": "yes", "dims": {}},
        {"text_id": "c", "text": "t3", "platform": "weibo", "brand": "大疆",
         "sentiment": "positive", "relevant": "no", "dims": {}},
    ]
    secondary = [
        {"text_id": "a", "text": "t1", "platform": "weibo", "brand": "大疆",
         "sentiment": "positive", "relevant": "yes", "dims": {"功能效果": "positive"}},
        {"text_id": "b", "text": "t2", "platform": "bilibili", "brand": "影石",
         "sentiment": "negative", "relevant": "no", "dims": {"价格价值": "negative"}},
        {"text_id": "c", "text": "t3", "platform": "weibo", "brand": "大疆",
         "sentiment": "positive", "relevant": "no", "dims": {"价格价值": "negative"}},
    ]
    out = tmp / "review.xlsx"
    n = cs.gen_dispute_review(primary, secondary, out, "digital3c")
    assert n == 2  # a: 情感分歧；b: 相关分歧；c: 仅维度分歧，不进入
    # 模拟人工填写判定列后再回填解析
    from openpyxl import load_workbook

    wb = load_workbook(out)
    ws = wb.worksheets[0]
    hdr = [str(c.value) for c in ws[1]]
    col_sent = hdr.index("人工判定情感") + 1
    col_rel = hdr.index("人工判定相关") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=col_sent).value = "neutral"
        ws.cell(row=r, column=col_rel).value = "yes"
    wb.save(out)
    dec = cs.load_human_decisions(out)
    assert set(dec) == {"a", "b"}
    print("✓ 分歧复核表生成/回填（仅情感/相关分歧）通过")


def test_load_excluded_texts() -> None:
    """--exclude-golden：读取裁判集原文集合，供新卷切分排除旧样本。"""
    tmp = Path(tempfile.mkdtemp(prefix="sms_excl_"))
    p = tmp / "golden.csv"
    p.write_text("text_id,text,sentiment\nh1,重复文本一,positive\nh2,重复文本二,neutral\n",
                 encoding="utf-8-sig")
    ex = cs.load_excluded_texts([p])
    assert ex == {"重复文本一", "重复文本二"}
    assert cs.load_excluded_texts([tmp / "missing.csv"]) == set()
    print("✓ --exclude-golden 排除文本集合 通过")


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    test_holdout_split_quotas_and_time_window()
    test_holdout_split_shortfall_clamp()
    test_holdout_split_missing_dates()
    test_holdout_freeze_and_use_guard()
    test_gold_revised_filter()
    test_acceptance_review_sheet_roundtrip()
    test_load_human_decisions()
    test_load_human_decisions_shared_text_id()
    test_gen_dispute_review_roundtrip()
    test_load_excluded_texts()
    print("验证纪律工程测试全部通过 ✅")


if __name__ == "__main__":
    main()
